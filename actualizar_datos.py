import os
import sys

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

import openpyxl
import json
from collections import Counter, defaultdict
import pandas as pd
import io
import re

def load_workbook_shared(path, data_only=True):
    """Abre archivos Excel de forma compartida, incluso si el usuario los tiene abiertos en Excel."""
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except Exception:
        try:
            import win32file, win32con
            handle = win32file.CreateFile(
                path,
                win32con.GENERIC_READ,
                win32con.FILE_SHARE_READ | win32con.FILE_SHARE_WRITE | win32con.FILE_SHARE_DELETE,
                None,
                win32con.OPEN_EXISTING,
                0,
                None
            )
            size = win32file.GetFileSize(handle)
            _, data = win32file.ReadFile(handle, size)
            win32file.CloseHandle(handle)
            return openpyxl.load_workbook(io.BytesIO(data), data_only=data_only)
        except Exception as e:
            print(f"   [ERROR] No se pudo leer {path}: {e}")
            raise e

import datetime

# Calendario Oficial de Feriados en Chile (2025, 2026, 2027)
FERIADOS_CHILE = {
    # 2025
    datetime.date(2025, 1, 1),   # Año Nuevo
    datetime.date(2025, 4, 18),  # Viernes Santo
    datetime.date(2025, 4, 19),  # Sábado Santo
    datetime.date(2025, 5, 1),   # Día del Trabajo
    datetime.date(2025, 5, 21),  # Glorias Navales
    datetime.date(2025, 6, 20),  # Pueblos Indígenas
    datetime.date(2025, 6, 29),  # San Pedro y San Pablo
    datetime.date(2025, 7, 16),  # Virgen del Carmen
    datetime.date(2025, 8, 15),  # Asunción de la Virgen
    datetime.date(2025, 9, 18),  # Fiestas Patrias
    datetime.date(2025, 9, 19),  # Glorias del Ejército
    datetime.date(2025, 10, 12), # Encuentro de Dos Mundos
    datetime.date(2025, 10, 31), # Iglesias Evangélicas
    datetime.date(2025, 11, 1),  # Todos los Santos
    datetime.date(2025, 12, 8),  # Inmaculada Concepción
    datetime.date(2025, 12, 25), # Navidad

    # 2026
    datetime.date(2026, 1, 1),   # Año Nuevo
    datetime.date(2026, 4, 3),   # Viernes Santo
    datetime.date(2026, 4, 4),   # Sábado Santo
    datetime.date(2026, 5, 1),   # Día del Trabajo
    datetime.date(2026, 5, 21),  # Glorias Navales
    datetime.date(2026, 6, 21),  # Pueblos Indígenas
    datetime.date(2026, 6, 29),  # San Pedro y San Pablo
    datetime.date(2026, 7, 16),  # Virgen del Carmen
    datetime.date(2026, 8, 15),  # Asunción de la Virgen
    datetime.date(2026, 9, 18),  # Fiestas Patrias
    datetime.date(2026, 9, 19),  # Glorias del Ejército
    datetime.date(2026, 10, 12), # Encuentro de Dos Mundos
    datetime.date(2026, 10, 31), # Iglesias Evangélicas
    datetime.date(2026, 11, 1),  # Todos los Santos
    datetime.date(2026, 12, 8),  # Inmaculada Concepción
    datetime.date(2026, 12, 25), # Navidad

    # 2027
    datetime.date(2027, 1, 1),   # Año Nuevo
    datetime.date(2027, 3, 26),  # Viernes Santo
    datetime.date(2027, 3, 27),  # Sábado Santo
    datetime.date(2027, 5, 1),   # Día del Trabajo
    datetime.date(2027, 5, 21),  # Glorias Navales
    datetime.date(2027, 6, 21),  # Pueblos Indígenas
    datetime.date(2027, 6, 28),  # San Pedro y San Pablo
    datetime.date(2027, 7, 16),  # Virgen del Carmen
    datetime.date(2027, 8, 15),  # Asunción de la Virgen
    datetime.date(2027, 9, 18),  # Fiestas Patrias
    datetime.date(2027, 9, 19),  # Glorias del Ejército
    datetime.date(2027, 10, 11), # Encuentro de Dos Mundos
    datetime.date(2027, 10, 31), # Iglesias Evangélicas
    datetime.date(2027, 11, 1),  # Todos los Santos
    datetime.date(2027, 12, 8),  # Inmaculada Concepción
    datetime.date(2027, 12, 25), # Navidad
}

def parse_comex_date(val):
    """Parsea fechas de múltiples formatos (datetime, string ISO, DD-MM-YYYY, DD/MM/YYYY, o serial de Excel)."""
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.date() if isinstance(val, datetime.datetime) else val
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ('none', 'nan', '#n/a', '#ref!', '01-01-1900', '1900-01-01', '00:00:00', 'no facturado'):
        return None
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y/%m/%d', '%Y-%m-%d %H:%M:%S', '%d-%m-%Y %H:%M:%S'):
        try:
            return datetime.datetime.strptime(val_str[:10], fmt[:8] if len(val_str) == 8 else fmt).date()
        except Exception:
            pass
    try:
        num = float(val_str)
        if 35000 < num < 60000:
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(num))).date()
    except Exception:
        pass
    return None

def add_business_days(start_date, n_days, holidays=FERIADOS_CHILE):
    """Suma n días hábiles excluyendo sábados, domingos y feriados chilenos."""
    if not start_date or n_days <= 0:
        return start_date
    curr = start_date
    added = 0
    while added < n_days:
        curr += datetime.timedelta(days=1)
        if curr.weekday() < 5 and curr not in holidays:
            added += 1
    return curr

def count_business_days(start_date, end_date, holidays=FERIADOS_CHILE):
    """Calcula los días hábiles transcurridos entre start_date y end_date (start excluido, end incluido)."""
    if not start_date or not end_date or end_date <= start_date:
        return 0
    curr = start_date + datetime.timedelta(days=1)
    count = 0
    while curr <= end_date:
        if curr.weekday() < 5 and curr not in holidays:
            count += 1
        curr += datetime.timedelta(days=1)
    return count

CUSTOMER_RULES = []

def load_customer_rules(base_dir=None):
    """Carga la matriz de reglas de negocio COMEX por cliente desde reglas_clientes.json."""
    global CUSTOMER_RULES
    if not base_dir:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    r_path = os.path.join(base_dir, "reglas_clientes.json")
    if os.path.exists(r_path):
        try:
            with open(r_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                CUSTOMER_RULES = data.get('reglas', [])
                print(f"-> [OK] {len(CUSTOMER_RULES)} reglas personalizadas de clientes cargadas desde reglas_clientes.json")
        except Exception as e:
            print("   [AVISO] Error leyendo reglas_clientes.json:", e)

def get_customer_rule(cli_name, cod_sap=None):
    """Retorna la regla COMEX personalizada para un cliente específico por código SAP o coincidencia de nombre."""
    if not cli_name and not cod_sap:
        return None
    cli_up = str(cli_name or '').strip().upper()
    cod_str = str(cod_sap or '').strip().lstrip('0')
    for r in CUSTOMER_RULES:
        r_cod = str(r.get('codigo_sap', '')).strip().lstrip('0')
        if cod_str and r_cod and r_cod == cod_str:
            return r
        pat = str(r.get('cliente_patron', '')).strip().upper()
        if pat and pat in cli_up:
            return r
    return None

def clean_client_name(c_name):
    """Limpia el nombre del cliente eliminando notas operativas, sufijos de pedidos previos (ex pedido 40xxxxxx) o redestinaciones."""
    if not c_name:
        return ""
    s = str(c_name).strip()
    s = re.sub(r'\s*\((?:intercompany|cliente interno)\)', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\(\[\{]?\s*(?:ex\s*pedido|ex-pedido|expedido|ex\s*ped|ex\d{6,}|anterior)\s*[:#]?\s*[\d\s\/\-_]+[\)\]\}]?', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'[\(\[\{]?\s*ex\s*\d{6,}[\s\d\/\-_]*[\)\]\}]?', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'\s*se\s+redestinar[aá].*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'\(\s*\)', '', s)
    s = re.sub(r'\[\s*\]', '', s)
    return re.sub(r'\s{2,}', ' ', s).strip()

def normalize_sap_doc(val):
    """Estandariza folios SAP, facturas y documentos contables eliminando ceros a la izquierda, decimales (.0) y espacios."""
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        if pd.isna(val) or val == 0:
            return ""
        val = str(int(val))
    s = str(val).strip()
    if not s or s.lower() in ('none', 'nan', 'null', '0', 'no facturado', '-', 'pend', 'pendiente', '#ref!', '#n/a'):
        return ""
    if s.endswith('.0'):
        s = s[:-2]
    s = s.lstrip('0').strip()
    return s

def evaluate_order_milestones(zarpe, draft_r, fin_r, fac, bl, dhl=None, swb=None, fin_e=None, today_dt=None, holidays=FERIADOS_CHILE, condicion_venta=None, nro_dhl_explicit=None, cli=None, co_aprobado=None, formato_liberacion=None):
    """Calcula hitos operativos: Estatus documental, BL Naviera (48h), Borrador Draft (4d), Finales (15d), DHL/Release y Reglas de Cliente."""
    if today_dt is None:
        today_dt = datetime.date.today()
    z_dt = parse_comex_date(zarpe)
    draft_r_dt = parse_comex_date(draft_r)
    fin_r_dt = parse_comex_date(fin_r)
    fac_str = normalize_sap_doc(fac)
    bl_str = str(bl or '').strip()
    has_fac = bool(fac_str)
    has_bl = bool(bl_str and bl_str.lower() not in ('none', 'nan', '', 'no emitido'))
    
    # 0. Consultar Reglas Personalizadas del Cliente (ej. Andes Asia, AquaChile Inc)
    c_rule = get_customer_rule(cli)
    categoria_cliente = c_rule.get('categoria', 'Comercial') if c_rule else 'Comercial'
    if not formato_liberacion and c_rule and c_rule.get('formato_liberacion'):
        formato_liberacion = c_rule.get('formato_liberacion')
    if not condicion_venta and c_rule and c_rule.get('condicion_venta'):
        condicion_venta = c_rule.get('condicion_venta')
    if str(condicion_venta or '').strip().lower() == 'intercompany':
        condicion_venta = 'Crédito'

    # 1. Procesar N° DHL y Fecha Envío DHL
    dhl_raw = str(dhl or '').strip()
    swb_raw = str(swb or '').strip()
    dhl_up = dhl_raw.upper()
    
    # Si la regla de cliente declara que NO requiere DHL:
    if c_rule and c_rule.get('requiere_dhl') is False:
        nro_dhl = "No aplica"
        fecha_dhl = "No Aplica"
    elif dhl_up in ('NA', 'N/A') or any(k in dhl_up for k in ('NO REQUIER', 'NO REPARTE', 'NO APLICA', 'NOAPLICA', 'CLIENTE RETIRA', 'NO DHL', 'SIN DHL', 'NO SE ENVIA')):
        nro_dhl = "No aplica"
        fecha_dhl = "No Aplica"
    elif dhl_raw:
        f_dhl_dt = parse_comex_date(dhl_raw)
        if f_dhl_dt:
            nro_dhl = str(nro_dhl_explicit or '').strip()
            fecha_dhl = f_dhl_dt.strftime('%Y-%m-%d')
        else:
            nro_dhl = dhl_raw
            fecha_dhl = ""
    else:
        nro_dhl = str(nro_dhl_explicit or '').strip()
        fecha_dhl = ""

    # Deducir formato de liberación si no viene explícito
    swb_low = str(swb_raw or '').lower()
    if not formato_liberacion or formato_liberacion == '-':
        if 'telex' in swb_low:
            formato_liberacion = 'Telex Release'
        elif any(k in swb_low for k in ('swb', 'waybill', 'express')):
            formato_liberacion = 'SWB'
        elif 'obl' in swb_low or 'original' in swb_low:
            formato_liberacion = 'OBL'
        elif 'destino' in swb_low:
            formato_liberacion = 'Liberación en Destino'
        elif c_rule and c_rule.get('formato_liberacion'):
            formato_liberacion = c_rule.get('formato_liberacion')
        elif categoria_cliente == 'Intercompany':
            formato_liberacion = 'SWB'
        else:
            # Clientes regulares / crédito: vienen por defecto como BL original (OBL)
            # no desde un inicio como SWB, tras unos días o semanas se solicita cambio a SWB, OBL o Telex
            formato_liberacion = 'OBL'
    
    # Limpiar cualquier remanente de '(Intercompany)' en formato
    formato_liberacion = re.sub(r'\s*\(intercompany\)', '', str(formato_liberacion or ''), flags=re.IGNORECASE).strip()
    if not formato_liberacion:
        formato_liberacion = 'SWB' if categoria_cliente == 'Intercompany' else 'OBL'

    # Indicadores operativos de despacho y liberación
    has_dhl_date = bool(fecha_dhl and fecha_dhl != 'No Aplica')
    has_dhl_guide = bool(nro_dhl and nro_dhl != 'No aplica')
    is_dhl_na = (nro_dhl == 'No aplica' or str(fecha_dhl).lower() == 'no aplica')

    # Solo fechas reales válidas en Release (eliminar 'cliente interno', 'SWB (Intercompany)', etc.)
    swb_dt_val = parse_comex_date(swb_raw)
    fecha_release_val = swb_dt_val.strftime('%Y-%m-%d') if swb_dt_val else ""
    has_release = bool(fecha_release_val or (swb_raw and any(k in swb_low for k in ('liberado', 'telex', 'swb enviado'))))

    is_anticipo = ('anticip' in str(condicion_venta or '').lower())
    has_co_aprobado = bool(parse_comex_date(co_aprobado) or (co_aprobado and str(co_aprobado).strip().lower() not in ('none', 'nan', '', '-', '0', 'crédito', 'credito')))

    # 2. Determinar Estatus
    if not has_fac:
        if z_dt:
            estatus = "Pendiente de facturar"
        else:
            estatus = "Por Despachar"
    elif c_rule and c_rule.get('criterio_cierre') == 'co_aprobado':
        # Regla Intercompany / Cierre por CO Aprobado (ej. Andes Asia, AquaChile Inc, Agrosuper):
        # Con factura y zarpe:
        # 1. Si el CO está aprobado -> Finalizado directo
        # 2. Si cuenta con Finales Real e hito de cierre o histórico -> Finalizado
        if has_co_aprobado or (fin_r_dt and (has_release or is_dhl_na or count_business_days(fin_r_dt, today_dt, holidays) >= 10)):
            estatus = "Finalizado"
        elif draft_r_dt and not fin_r_dt:
            estatus = "Finales en Trámite"
        else:
            estatus = "Solicitando CO"
    else:
        # Operación Facturada regular:
        # Regla Anticipo: Requiere Release Y (Fecha DHL o No Aplica o Guía)
        if is_anticipo:
            if has_release and (has_dhl_date or is_dhl_na or has_dhl_guide):
                estatus = "Finalizado"
            elif fin_r_dt:
                estatus = "Pendiente Liberación / Pago"
            elif draft_r_dt:
                estatus = "Finales en Trámite"
            else:
                estatus = "Solicitando CO"
        # Regla General / Crédito: con Finales Real y (Fecha DHL o No Aplica o Release o Guía)
        elif fin_r_dt and (has_dhl_date or is_dhl_na or has_release or has_dhl_guide):
            estatus = "Finalizado"
        # Cierre Histórico: Si Finales Real se completó hace más de 10 días hábiles sin pendientes
        elif fin_r_dt and count_business_days(fin_r_dt, today_dt, holidays) >= 10:
            estatus = "Finalizado"
        elif draft_r_dt and fin_r_dt:
            estatus = "Pendiente Despacho DHL / SWB"
        elif not draft_r_dt and not fin_r_dt:
            estatus = "Solicitando CO"
        elif draft_r_dt and not fin_r_dt:
            estatus = "Finales en Trámite"
        else:
            estatus = "En Trámite"

    # Regla: Si está "Pendiente de facturar" o "Por Despachar", draft y finales quedan vacíos porque aún no se factura
    if estatus in ("Pendiente de facturar", "Por Despachar"):
        bl_esperado_dt = add_business_days(z_dt, 2, holidays) if z_dt else None
        bl_esperado_str = bl_esperado_dt.strftime('%Y-%m-%d') if bl_esperado_dt else ''
        return {
            "zarpe": z_dt.strftime('%Y-%m-%d') if z_dt else (str(zarpe)[:10] if zarpe else ''),
            "estatus": estatus,
            "bl_esperado": bl_esperado_str,
            "estado_bl": "EN PLAZO NAVIERA" if (z_dt and today_dt <= bl_esperado_dt) else ("DEMORADO NAVIERA" if z_dt else "PENDIENTE ZARPE"),
            "alerta_bl": f"Naviera excede 48h hábiles sin emitir BL (venció {bl_esperado_str})" if (z_dt and today_dt > bl_esperado_dt) else "",
            "dias_atraso_bl": count_business_days(bl_esperado_dt, today_dt, holidays) if (z_dt and today_dt > bl_esperado_dt) else 0,
            "draft_esperado": "",
            "draft_real": "",
            "estado_draft": "",
            "dias_atraso_draft": 0,
            "finales_esperado": "",
            "finales_real": "",
            "estado_finales": "",
            "dias_atraso_finales": 0,
            "nro_dhl": nro_dhl,
            "fecha_dhl": fecha_dhl,
            "swb_liberacion": fecha_release_val,
            "formato_liberacion": formato_liberacion,
            "categoria_cliente": categoria_cliente
        }

    # 3. Hito BL Naviera (Tope 48h hábiles)
    bl_esperado_dt = add_business_days(z_dt, 2, holidays) if z_dt else None
    bl_esperado_str = bl_esperado_dt.strftime('%Y-%m-%d') if bl_esperado_dt else ''
    
    if has_bl or has_fac:
        estado_bl = "BL RECIBIDO / FACTURADO"
        alerta_bl = ""
        dias_atraso_bl = 0
    elif z_dt and bl_esperado_dt:
        if today_dt <= bl_esperado_dt:
            estado_bl = "EN PLAZO NAVIERA"
            alerta_bl = ""
            dias_atraso_bl = 0
        else:
            dias_atraso_bl = count_business_days(bl_esperado_dt, today_dt, holidays)
            estado_bl = f"DEMORADO NAVIERA (+{dias_atraso_bl}d)"
            alerta_bl = f"Naviera excede 48h hábiles sin emitir BL (venció {bl_esperado_str})"
    else:
        estado_bl = "PENDIENTE ZARPE"
        alerta_bl = ""
        dias_atraso_bl = 0
        
    # 4. Hito Borrador / Draft (Tope 4 días hábiles: 2d naviera + 2d agencia)
    draft_esperado_dt = add_business_days(z_dt, 4, holidays) if z_dt else None
    draft_esperado_str = draft_esperado_dt.strftime('%Y-%m-%d') if draft_esperado_dt else ''
    
    if draft_r_dt:
        draft_real_str = draft_r_dt.strftime('%Y-%m-%d')
        if draft_esperado_dt and draft_r_dt > draft_esperado_dt:
            st_draft = "ATRASADO"
            d_draft = count_business_days(draft_esperado_dt, draft_r_dt, holidays)
        else:
            st_draft = "EN PLAZO"
            d_draft = 0
    else:
        draft_real_str = ''
        if not z_dt:
            st_draft = "PENDIENTE ZARPE"
            d_draft = 0
        elif today_dt <= draft_esperado_dt:
            st_draft = "EN PLAZO"
            d_draft = 0
        else:
            st_draft = "ATRASADO"
            d_draft = count_business_days(draft_esperado_dt, today_dt, holidays)
            
    # 5. Hito Documentos Finales (Tope 15 días hábiles)
    fin_e_dt = add_business_days(z_dt, 15, holidays) if z_dt else parse_comex_date(fin_e)
    fin_esperado_str = fin_e_dt.strftime('%Y-%m-%d') if fin_e_dt else (str(fin_e)[:10] if fin_e else '')
    
    if fin_r_dt:
        fin_real_str = fin_r_dt.strftime('%Y-%m-%d')
        if fin_e_dt and fin_r_dt > fin_e_dt:
            st_fin = "ATRASADO"
            d_fin = count_business_days(fin_e_dt, fin_r_dt, holidays)
        else:
            st_fin = "EN PLAZO"
            d_fin = 0
    else:
        fin_real_str = ''
        if not z_dt or not fin_e_dt or today_dt <= fin_e_dt:
            st_fin = "EN PLAZO"
            d_fin = 0
        else:
            st_fin = "ATRASADO"
            d_fin = count_business_days(fin_e_dt, today_dt, holidays)
            
    return {
        "zarpe": z_dt.strftime('%Y-%m-%d') if z_dt else (str(zarpe)[:10] if zarpe else ''),
        "estatus": estatus,
        "bl_esperado": bl_esperado_str,
        "estado_bl": estado_bl,
        "alerta_bl": alerta_bl,
        "dias_atraso_bl": dias_atraso_bl,
        "draft_esperado": draft_esperado_str,
        "draft_real": draft_real_str,
        "estado_draft": st_draft,
        "dias_atraso_draft": d_draft,
        "finales_esperado": fin_esperado_str,
        "finales_real": fin_real_str,
        "estado_finales": st_fin,
        "dias_atraso_finales": d_fin,
        "nro_dhl": nro_dhl,
        "fecha_dhl": fecha_dhl,
        "swb_liberacion": fecha_release_val,
        "formato_liberacion": formato_liberacion,
        "categoria_cliente": categoria_cliente
    }

def main():
    print("==================================================================")
    print("   AQUASHIELD - CONTROL CONGELADO: CONSOLIDACION Y TRIANGULACION")
    print("==================================================================")
    
    base_dir = os.path.dirname(os.path.abspath(__file__))
    os.chdir(base_dir)

    # 0. Cargar Matriz de Reglas de Clientes COMEX
    load_customer_rules(base_dir)

    # 1. Definición Maestra de Equipos COMEX
    EQUIPOS_AUTORES = {
        # EQUIPO FRESCO (Omitir completamente de Control Congelado)
        'DPINCOL':    {'analista': 'Daphne',          'equipo': 'Fresco',    'apoyo': False},
        'SMOHOR':     {'analista': 'Scandar',         'equipo': 'Fresco',    'apoyo': False},
        'VVENEGAS':   {'analista': 'Victor',          'equipo': 'Fresco',    'apoyo': False},
        'ESILVAJ':    {'analista': 'Evelyn',          'equipo': 'Fresco',    'apoyo': False},
        'CZAMORANO':  {'analista': 'Cristina',        'equipo': 'Fresco',    'apoyo': False},

        # EQUIPO CONGELADO (Permanecen en Control Congelado)
        'JOSANCHEZM': {'analista': 'Jonathan',        'equipo': 'Congelado', 'apoyo': True},
        'EMARTINEZG': {'analista': 'Eugenia',         'equipo': 'Congelado', 'apoyo': True},
        'EMARTINEZ':  {'analista': 'Eugenia',         'equipo': 'Congelado', 'apoyo': True},
        'YVALERO':    {'analista': 'Andreina',        'equipo': 'Congelado', 'apoyo': True},
        'MIFIGUEROA': {'analista': 'Matías Figueroa', 'equipo': 'Congelado', 'apoyo': True},
        'MNSOTO':     {'analista': 'Matías Soto',     'equipo': 'Congelado', 'apoyo': True},
        'LCONEJEROSM':{'analista': 'Libni',           'equipo': 'Congelado', 'apoyo': True},
        'SNAGUIL':    {'analista': 'Soraya',          'equipo': 'Congelado', 'apoyo': True},
        'MRAMIREZ':   {'analista': 'Andreina',        'equipo': 'Congelado', 'apoyo': True} # Reemplazo cumplido -> Andreina
    }

    fresco_authors_set = {k for k, v in EQUIPOS_AUTORES.items() if v['equipo'] == 'Fresco'}
    congelado_authors_set = {k for k, v in EQUIPOS_AUTORES.items() if v['equipo'] == 'Congelado'}

    # 2. Cargar Planilla VF05N (Facturas de SAP y sus Autores)
    factura_author = {}
    factura_date = {}

    vf_candidates = ["vf05n.xlsx", "transaccion_vf05n.xlsx", "VF05N.xlsx", "facturas_vf05n.xlsx"]
    vf_path = next((os.path.join(base_dir, f) for f in vf_candidates if os.path.exists(os.path.join(base_dir, f))), None)

    if vf_path:
        print(f"-> Procesando transaccion VF05N ({os.path.basename(vf_path)})...")
        wb_vf = load_workbook_shared(vf_path)
        ws_vf = wb_vf['Data'] if 'Data' in wb_vf.sheetnames else wb_vf.active

        # Detectar columnas si es tabla de autores o exportación de facturas
        col_header_1 = str(ws_vf.cell(row=1, column=1).value or '').strip().lower()
        col_header_2 = str(ws_vf.cell(row=1, column=2).value or '').strip().lower()
        col_header_3 = str(ws_vf.cell(row=1, column=3).value or '').strip().lower()

        is_mapping_table = ('autor' in col_header_1 and 'equipo' in col_header_3)

        if not is_mapping_table:
            for r in range(2, ws_vf.max_row + 1):
                doc = ws_vf.cell(row=r, column=1).value
                f_date = ws_vf.cell(row=r, column=2).value
                autor = ws_vf.cell(row=r, column=3).value
                if doc and autor:
                    doc_clean = normalize_sap_doc(doc)
                    if doc_clean:
                        factura_author[doc_clean] = str(autor).strip().upper()
                        if f_date:
                            f_date_clean = parse_comex_date(f_date)
                            factura_date[doc_clean] = f_date_clean.strftime('%Y-%m-%d') if f_date_clean else str(f_date)[:10]
            print(f"   [OK] {len(factura_author)} facturas con autor SAP registradas desde VF05N.")
        else:
            print("   [INFO] El archivo contiene la tabla de mapeo de equipos.")
            # Si en el archivo hay mapeo de equipos, actualizamos diccionario
            for r in range(2, ws_vf.max_row + 1):
                aut = str(ws_vf.cell(row=r, column=1).value or '').strip().upper()
                ana = str(ws_vf.cell(row=r, column=2).value or '').strip()
                eq = str(ws_vf.cell(row=r, column=3).value or '').strip()
                if aut and eq:
                    es_congelado = ('congelado' in eq.lower())
                    EQUIPOS_AUTORES[aut] = {
                        'analista': ana,
                        'equipo': 'Congelado' if es_congelado else 'Fresco',
                        'apoyo': es_congelado
                    }
    else:
        print("   [AVISO] No se encontró archivo VF05N. Usando histórico.")

    def clean_analyst(name):
        if not name: return "Sin Asignar"
        n = str(name).strip().upper()
        if 'EUGENIA' in n: return 'Eugenia'
        if 'ANDREINA' in n: return 'Andreina'
        if 'JONATHAN' in n or 'JONA' in n: return 'Jonathan'
        if 'FIGUEROA' in n: return 'Matías Figueroa'
        if 'SOTO' in n: return 'Matías Soto'
        if 'LIBNI' in n: return 'Libni'
        if 'SORAYA' in n: return 'Soraya'
        if 'MARCELO' in n: return 'Andreina'
        if 'TERRESTRE' in n: return 'Terrestre'
        if any(f in n for f in ['SCANDAR', 'CRISTINA', 'DAPHNE', 'VICTOR', 'EVELYN', 'FRESCO']):
            return 'FRESCO_OMITIR'
        # Salvaguarda: si el texto no corresponde a un analista válido, no es analista
        return "Sin Asignar"

    # 3. Cargar Asignación de Clientes (Cartera oficial)
    cliente_portfolio = {}
    codigo_portfolio = {}
    cliente_codigo_map = {}
    cartera_raw_list = []
    asig_path = os.path.join(base_dir, "Asignacion Clientes.xlsx")
    if os.path.exists(asig_path):
        print("-> Procesando Asignacion Clientes.xlsx...")
        wb_asig = load_workbook_shared(asig_path)
        ws_asig = wb_asig['Asignación'] if 'Asignación' in wb_asig.sheetnames else (wb_asig['Asignacion Clientes'] if 'Asignacion Clientes' in wb_asig.sheetnames else wb_asig.active)
        
        # Detectar columnas: Col 1 = Código numérico, Col 2 = Nombre cliente texto, Col 3 = Asignación analista
        col_resp_idx = None
        for c in range(1, ws_asig.max_column + 1):
            h_val = str(ws_asig.cell(row=1, column=c).value or '').strip().lower()
            if 'asignaci' in h_val or 'responsable' in h_val or 'analista' in h_val:
                col_resp_idx = c
                break
        if not col_resp_idx:
            col_resp_idx = 3 if ws_asig.max_column >= 3 else 2
            
        col_name_idx = 2 if col_resp_idx == 3 else 1
        col_code_idx = 1 if col_resp_idx == 3 else None

        for r in range(2, ws_asig.max_row + 1):
            raw_resp = ws_asig.cell(row=r, column=col_resp_idx).value
            raw_name = ws_asig.cell(row=r, column=col_name_idx).value
            raw_code = ws_asig.cell(row=r, column=col_code_idx).value if col_code_idx else None
            raw_cont = ws_asig.cell(row=r, column=5).value if ws_asig.max_column >= 5 else None
            raw_cel = ws_asig.cell(row=r, column=6).value if ws_asig.max_column >= 6 else None
            
            if raw_resp:
                ana = clean_analyst(str(raw_resp).strip())
                if ana and ana != 'Sin Asignar':
                    code_clean = str(raw_code).strip() if raw_code else ''
                    name_clean = str(raw_name).strip() if raw_name else ''
                    if name_clean:
                        c_clean = name_clean.upper().replace('"', '').replace('«', '').replace('»', '').strip()
                        if c_clean:
                            cliente_portfolio[c_clean] = ana
                            cartera_raw_list.append({
                                "codigo": code_clean,
                                "cliente": name_clean,
                                "analista": ana,
                                "contacto": str(raw_cont or '').strip(),
                                "celular": str(raw_cel or '').strip()
                            })
                            if code_clean:
                                cliente_codigo_map[c_clean] = code_clean
                                c_alpha = re.sub(r'[^\w\s]', '', c_clean).strip()
                                cliente_codigo_map[c_alpha] = code_clean
                    if code_clean:
                        codigo_portfolio[code_clean] = ana
                            
        print(f"   [OK] {len(cliente_portfolio)} clientes en cartera oficial cargados con sus analistas correspondientes ({len(cliente_codigo_map)} códigos mapeados).")
    else:
        print("   [AVISO] No se encontro Asignacion Clientes.xlsx")

    def find_sap_code_for_client(c_name):
        if not c_name:
            return ""
        c_up = str(c_name).strip().upper()
        if c_up in cliente_codigo_map:
            return cliente_codigo_map[c_up]
        
        # Consultar si tiene regla explícita
        rule = get_customer_rule(c_name)
        if rule and rule.get("codigo_sap"):
            return str(rule["codigo_sap"]).strip()

        # Limpiar notas operativas, números de pedido previos o anotaciones
        c_clean = re.sub(r'[\(\[\{]?\s*EX\s*PEDIDO\s*[\d\s\/\-_]+[\)\]\}]?', '', c_up, flags=re.IGNORECASE).strip()
        c_clean = re.sub(r'[\(\[\{]?\s*EX\s*[\d\s\/\-_]+[\)\]\}]?', '', c_clean, flags=re.IGNORECASE).strip()
        c_clean = re.sub(r'[\(\[\{]?\s*ANTERIOR\s*[\d\s\/\-_]+[\)\]\}]?', '', c_clean, flags=re.IGNORECASE).strip()
        c_clean = re.sub(r'\b(SE REDESTINAR[AÁ]|TAILANDIA|VIETNAM)\b', '', c_clean, flags=re.IGNORECASE).strip()
        c_clean = c_clean.replace('"', '').replace('«', '').replace('»', '').replace('`', '').strip()
        
        if c_clean in cliente_codigo_map:
            return cliente_codigo_map[c_clean]
            
        c_alpha = re.sub(r'[^\w\s]', '', c_clean).strip()
        if c_alpha in cliente_codigo_map:
            return cliente_codigo_map[c_alpha]
            
        # Normalizar sufijos societarios típicos
        c_norm = re.sub(r'\b(CO|LTD|LLC|LIMITED|INC|CORP|SA|SRL|PT|SC|OOO|SP Z O O)\b', '', c_alpha).strip()
        c_norm = re.sub(r'\s+', ' ', c_norm)
        
        for item in cartera_raw_list:
            k = item['cliente'].upper().replace('"', '').replace('«', '').replace('»', '').strip()
            k_alpha = re.sub(r'[^\w\s]', '', k).strip()
            k_norm = re.sub(r'\b(CO|LTD|LLC|LIMITED|INC|CORP|SA|SRL|PT|SC|OOO|SP Z O O)\b', '', k_alpha).strip()
            k_norm = re.sub(r'\s+', ' ', k_norm)
            
            if c_alpha == k_alpha:
                return item.get('codigo', '')
            if len(c_norm) > 4 and len(k_norm) > 4:
                if c_norm == k_norm or c_norm in k_norm or k_norm in c_norm:
                    return item.get('codigo', '')
            if len(c_alpha) > 5 and len(k_alpha) > 5:
                if c_alpha in k_alpha or k_alpha in c_alpha:
                    return item.get('codigo', '')
                    
        return ""

    # 4. Cargar Control Congelado (Activos en curso)
    cc_file = "Control Congelado Comex 2026.xlsx"
    if not os.path.exists(cc_file):
        cc_file = "Control Congelado Comex 2026 - Optimizado.xlsx"
        
    print(f"-> Procesando {cc_file}...")
    wb_cc = load_workbook_shared(cc_file)
    ws_cc = wb_cc['CONTROL CONGELADO'] if 'CONTROL CONGELADO' in wb_cc.sheetnames else wb_cc.active

    # Cargar y combinar feriados adicionales si existen en la planilla
    feriados_set = set(FERIADOS_CHILE)
    if 'Feriados' in wb_cc.sheetnames:
        ws_fer = wb_cc['Feriados']
        for r_fer in range(1, ws_fer.max_row + 1):
            val_fer = ws_fer.cell(row=r_fer, column=1).value
            f_dt = parse_comex_date(val_fer)
            if f_dt and f_dt.year in (2025, 2026, 2027):
                feriados_set.add(f_dt)
    today_dt = datetime.date.today()
    print(f"   [OK] Calendario operativo: {len(feriados_set)} feriados registrados (2025-2027). Fecha hoy: {today_dt}")

    # Saneamiento de visibilidad: asegurar que la planilla base no contenga filas ocultas ni filtros pre-activos
    has_hidden = any(dim.hidden for dim in ws_cc.row_dimensions.values())
    if has_hidden:
        for dim in ws_cc.row_dimensions.values():
            dim.hidden = False
        try:
            from openpyxl.utils import get_column_letter
            ws_cc.auto_filter.ref = f"A1:{get_column_letter(ws_cc.max_column)}{ws_cc.max_row}"
            if hasattr(ws_cc, 'sheet_view') and ws_cc.sheet_view:
                ws_cc.sheet_view.selection[0].activeCell = 'A2'
                ws_cc.sheet_view.selection[0].sqref = 'A2'
            wb_cc.save(cc_file)
            print(f"   [OK] Saneamiento Excel: 100% filas desocultadas y filtro restablecido en {cc_file}.")
        except Exception as e_save:
            print("   [AVISO] No se pudo re-guardar visibilidad en Excel:", e_save)

    active_shipments = []
    history_client_resp = {}
    excel_ped_resp = {}  # Mapeo exacto de pedido a responsable desde la planilla oficial
    activos_fresco_omitidos = 0

    # Cargar previamente cambios manuales si existen para enriquecer excel_ped_resp
    json_manual_path = "cambios_manuales.json"
    if os.path.exists(json_manual_path):
        try:
            with open(json_manual_path, "r", encoding="utf-8") as f_man:
                man_data = json.load(f_man)
                for p_str, ed in man_data.get("edits", {}).items():
                    if ed.get("responsable"):
                        excel_ped_resp[str(p_str).strip()] = clean_analyst(ed.get("responsable"))
        except Exception:
            pass

    # Mapeo dinámico e inteligente de columnas según nombres de cabecera
    header_col = {}
    for c in range(1, ws_cc.max_column + 1):
        h_val = str(ws_cc.cell(row=1, column=c).value or '').strip().upper()
        if 'RESPONSABLE' in h_val: header_col['resp'] = c
        elif 'SOLICITANTE' in h_val or 'CLIENTE' in h_val: header_col['cli'] = c
        elif 'PEDIDO' in h_val: header_col['ped'] = c
        elif 'NAVE' in h_val: header_col['nave'] = c
        elif 'ZARPE' in h_val or 'P44' in h_val: header_col['zarpe'] = c
        elif 'ETA' in h_val: header_col['eta'] = c
        elif 'PUERTO' in h_val: header_col['puerto'] = c
        elif 'FOLIO FACTURA' in h_val or h_val == 'FACTURA': header_col['fac'] = c
        elif 'FECHA FACTURA' in h_val or 'F. FACTURA' in h_val: header_col['f_fac'] = c
        elif 'REAL' in h_val and 'DRAFT' in h_val: header_col['draft_r'] = c
        elif 'ESPERADO' in h_val and 'DRAFT' in h_val: header_col['draft_e'] = c
        elif 'PLAZO' in h_val and 'DRAFT' in h_val: header_col['draft_p'] = c
        elif 'CERTIFICADO' in h_val: header_col['cert'] = c
        elif 'FINALES' in h_val and ('MÁX' in h_val or 'MAX' in h_val or 'REAL' in h_val): header_col['fin_r'] = c
        elif 'FINALES' in h_val and 'ESPERADO' in h_val: header_col['fin_e'] = c
        elif 'POST VENTA' in h_val or 'MACA' in h_val: header_col['pv'] = c
        elif 'F. DHL' in h_val or 'FECHA DHL' in h_val or 'F.DHL' in h_val: header_col['f_dhl'] = c
        elif 'GUIA DHL' in h_val or 'GUÍA DHL' in h_val or 'NRO DHL' in h_val or 'N° DHL' in h_val: header_col['nro_dhl'] = c
        elif 'DHL' in h_val: header_col['dhl'] = c
        elif 'SWB' in h_val or 'RELEASE' in h_val: header_col['swb'] = c
        elif 'CONDICION' in h_val or 'CONDICIÓN' in h_val or 'VENTA' in h_val: header_col['condicion'] = c
        elif 'OBSERVACI' in h_val: header_col['obs'] = c

    for r in range(2, ws_cc.max_row + 1):
        resp = ws_cc.cell(row=r, column=header_col.get('resp', 1)).value
        cli = ws_cc.cell(row=r, column=header_col.get('cli', 2)).value
        ped = ws_cc.cell(row=r, column=header_col.get('ped', 3)).value
        nave = ws_cc.cell(row=r, column=header_col.get('nave', 4)).value
        zarpe = ws_cc.cell(row=r, column=header_col.get('zarpe', 5)).value
        eta = ws_cc.cell(row=r, column=header_col.get('eta', 6)).value
        puerto = ws_cc.cell(row=r, column=header_col.get('puerto', 7)).value
        fac = ws_cc.cell(row=r, column=header_col.get('fac', 8)).value
        f_fac = ws_cc.cell(row=r, column=header_col['f_fac']).value if 'f_fac' in header_col else None
        draft_r = ws_cc.cell(row=r, column=header_col.get('draft_r', 10)).value
        draft_p = ws_cc.cell(row=r, column=header_col.get('draft_p', 11)).value
        cert = ws_cc.cell(row=r, column=header_col.get('cert', 12)).value
        fin_r = ws_cc.cell(row=r, column=header_col.get('fin_r', 13)).value
        fin_e = ws_cc.cell(row=r, column=header_col.get('fin_e', 14)).value
        pv = ws_cc.cell(row=r, column=header_col.get('pv', 15)).value
        dhl_cell = ws_cc.cell(row=r, column=header_col.get('f_dhl', header_col.get('dhl', 16))).value
        guia_cell = ws_cc.cell(row=r, column=header_col['nro_dhl']).value if 'nro_dhl' in header_col else None
        swb = ws_cc.cell(row=r, column=header_col.get('swb', 17)).value
        obs = ws_cc.cell(row=r, column=header_col.get('obs', 18)).value
        
        if ped or resp:
            ped_str = str(ped).strip().lstrip('0')
            fac_str = normalize_sap_doc(fac)

            # REGLA DE EXCLUSION FRESCO: Si la factura fue emitida por el equipo Fresco, omitir
            if fac_str and fac_str in factura_author:
                if factura_author[fac_str] in fresco_authors_set:
                    activos_fresco_omitidos += 1
                    continue

            analyst = clean_analyst(resp)
            if analyst == 'FRESCO_OMITIR':
                activos_fresco_omitidos += 1
                continue

            if ped_str:
                excel_ped_resp[ped_str] = analyst

            if cli:
                cli_clean = clean_client_name(cli)
                history_client_resp[str(cli_clean).strip().upper()] = analyst
                
            # REGLA DE CORTE HISTÓRICO: A partir de Junio 2026 comienza la planilla oficial
            z_dt = parse_comex_date(zarpe)
            if z_dt and z_dt < datetime.date(2026, 6, 1):
                continue

            # Condición de Venta (Crédito, Anticipo, etc.)
            cond_venta = ''
            if 'condicion' in header_col:
                c_cell = ws_cc.cell(row=r, column=header_col['condicion']).value
                if c_cell:
                    cond_venta = str(c_cell).strip()
            if not cond_venta:
                if cert and 'cred' in str(cert).strip().lower():
                    cond_venta = 'Crédito'
                elif swb and 'cred' in str(swb).strip().lower():
                    cond_venta = 'Crédito'
                elif obs and 'anticip' in str(obs).strip().lower():
                    cond_venta = 'Anticipo'
                elif obs and 'cred' in str(obs).strip().lower():
                    cond_venta = 'Crédito'

            # Limpieza y estandarización de fechas y textos
            eta_dt = parse_comex_date(eta)
            eta_fmt = eta_dt.strftime('%Y-%m-%d') if eta_dt else ''

            cert_dt = parse_comex_date(cert)
            if cert_dt:
                cert_val = cert_dt.strftime('%Y-%m-%d')
            elif cert and str(cert).strip().lower() not in ('none', 'nan', 'crédito', 'credito'):
                cert_val = str(cert).strip()
            else:
                cert_val = ''

            # SWB / Release: Si es fecha, guardar ISO. Si es texto (ej: "cliente interno", "OBL"), deducir formato y limpiar fecha
            swb_dt = parse_comex_date(swb)
            swb_val = swb_dt.strftime('%Y-%m-%d') if swb_dt else ''

            fmt_from_excel = None
            if swb and not swb_dt:
                s_low = str(swb).lower()
                if 'telex' in s_low:
                    fmt_from_excel = 'Telex Release'
                elif 'swb' in s_low or 'interno' in s_low:
                    fmt_from_excel = 'SWB'
                elif 'obl' in s_low or 'original' in s_low:
                    fmt_from_excel = 'OBL'
                elif 'destino' in s_low:
                    fmt_from_excel = 'Liberación en Destino'

            # Evaluación precisa de hitos operativos (48h BL naviera, 4d hábiles draft, 15d finales, DHL, Estatus y Reglas de Cliente)
            m = evaluate_order_milestones(zarpe, draft_r, fin_r, fac_str, None, dhl=dhl_cell, swb=swb_val, fin_e=fin_e, today_dt=today_dt, holidays=feriados_set, condicion_venta=cond_venta, nro_dhl_explicit=guia_cell, cli=cli, co_aprobado=cert_val, formato_liberacion=fmt_from_excel)
                
            author_code = factura_author.get(fac_str, '')
            author_sap_name = EQUIPOS_AUTORES.get(author_code, {}).get('analista', author_code)
            if author_sap_name and author_sap_name.lower() != analyst.lower():
                metodo_asign = f"Planilla Excel ({analyst}) [Apoyo facturación SAP: {author_sap_name}]"
            else:
                metodo_asign = f"Planilla Excel ({analyst})"

            f_fac_dt = parse_comex_date(f_fac)
            f_fac_fmt = f_fac_dt.strftime('%Y-%m-%d') if f_fac_dt else ''
            fecha_factura_val = f_fac_fmt if f_fac_fmt else factura_date.get(fac_str, '')

            final_swb = m["swb_liberacion"]

            active_shipments.append({
                "id": r - 1,
                "tipo": "Activo",
                "estatus": m["estatus"],
                "pedido": ped_str,
                "codigo_cliente": find_sap_code_for_client(clean_client_name(cli)),
                "responsable": analyst,
                "autor_sap": author_sap_name,
                "metodo_asignacion": metodo_asign,
                "cliente": clean_client_name(cli),
                "categoria_cliente": m["categoria_cliente"],
                "formato_liberacion": m["formato_liberacion"],
                "nave": str(nave or '').strip(),
                "zarpe": m["zarpe"],
                "eta": eta_fmt,
                "puerto": str(puerto or '').strip(),
                "condicion_venta": cond_venta,
                "factura": fac_str,
                "doc_factura": fac_str,
                "fecha_factura": fecha_factura_val,
                "bl_esperado": m["bl_esperado"],
                "estado_bl": m["estado_bl"],
                "alerta_bl": m["alerta_bl"],
                "dias_atraso_bl": m["dias_atraso_bl"],
                "draft_real": m["draft_real"],
                "draft_esperado": m["draft_esperado"],
                "draft_plazo": m["draft_esperado"] or "",
                "certificados": cert_val,
                "co_aprobado": cert_val,
                "finales_real": m["finales_real"],
                "finales_esperado": m["finales_esperado"],
                "nro_dhl": m["nro_dhl"],
                "fecha_dhl": m["fecha_dhl"],
                "dhl": m["nro_dhl"] or m["fecha_dhl"] or str(dhl_cell or '').strip(),
                "swb": final_swb,
                "swb_liberacion": final_swb,
                "release": final_swb,
                "post_venta": str(pv)[:10] if pv else '',
                "observacion": str(obs or '').strip(),
                "estado_draft": m["estado_draft"],
                "estado_finales": m["estado_finales"],
                "dias_atraso_draft": m["dias_atraso_draft"],
                "dias_atraso_finales": m["dias_atraso_finales"]
            })

    print(f"   [OK] {len(active_shipments)} despachos activos de Congelado conservados ({activos_fresco_omitidos} omitidos por ser de Fresco).")

    # Mapeo de Naves conocidas a Puerto de Salida (POL) para auto-completar pedidos de flujo/SQL
    nave_to_port = {}
    for item in active_shipments:
        nv = item.get("nave", "").strip().upper()
        pt = item.get("puerto", "").strip()
        if nv and pt and nv not in nave_to_port:
            nave_to_port[nv] = pt

    # 5. Cargar Flujo SAP (flujo v2 o v1)
    flujo_file = "flujo v2.xlsx" if os.path.exists("flujo v2.xlsx") else "flujo v1.xlsx"
    flow_orders = []
    flujo_fresco_omitidos = 0
    
    if os.path.exists(flujo_file):
        print(f"-> Procesando {flujo_file} (Flujo SAP por zarpar)...")
        wb_fl = load_workbook_shared(flujo_file)
        ws_fl = wb_fl['Data'] if 'Data' in wb_fl.sheetnames else wb_fl.active

        active_peds_set = set(s['pedido'] for s in active_shipments)

        def resolve_flow_analyst(cli_name, doc_fac, ped_str=None, cli_code=None):
            doc_clean = normalize_sap_doc(doc_fac)
            author_code = factura_author.get(doc_clean, '') if doc_clean else ''
            author_sap = EQUIPOS_AUTORES.get(author_code, {}).get('analista', author_code) if author_code else ''
            author_desc = author_sap or (author_code if author_code else "Sin Factura")

            # Regla 0: Prioridad Absoluta de la Planilla Excel / Reasignación Manual
            # Si el pedido tiene analista asignado en la planilla o fue reasignado, se respeta al 100%
            if ped_str and ped_str in excel_ped_resp:
                ana = excel_ped_resp[ped_str]
                if ana and ana != 'Sin Asignar':
                    metodo = f"Planilla Excel ({ana})"
                    if author_sap and author_sap.lower() != ana.lower():
                        metodo += f" [Apoyo facturación SAP: {author_sap}]"
                    return ana, author_desc, metodo

            # Regla de Exclusión Fresco: Si la factura fue emitida por el equipo Fresco, omitir
            if author_code in fresco_authors_set:
                return 'FRESCO_OMITIR', author_code, f"Facturado por Fresco ({author_sap})"

            # Regla 1: Cartera Asignada Oficial (por código SAP o nombre limpio)
            c_clean = clean_client_name(cli_name)
            c_code = str(cli_code or '').strip() or find_sap_code_for_client(c_clean)
            c_up = c_clean.strip().upper().replace('"', '').replace('«', '').replace('»', '') if c_clean else ''
            portfolio_ana = None

            if c_code and c_code in codigo_portfolio:
                portfolio_ana = codigo_portfolio[c_code]
            elif c_up:
                if c_up in cliente_portfolio:
                    portfolio_ana = cliente_portfolio[c_up]
                else:
                    for k, v in cliente_portfolio.items():
                        if (len(k) > 4 and k in c_up) or (len(c_up) > 4 and c_up in k):
                            portfolio_ana = v
                            break

            if portfolio_ana:
                if portfolio_ana == 'FRESCO_OMITIR':
                    return 'FRESCO_OMITIR', author_desc, "Cartera Fresco"
                if portfolio_ana != 'Sin Asignar':
                    metodo = f"Cartera Asignada ({portfolio_ana})"
                    if author_sap and author_sap.lower() != portfolio_ana.lower():
                        metodo += f" [Apoyo facturación SAP: {author_sap}]"
                    return portfolio_ana, author_desc, metodo

            # Regla 2: Historial Control Congelado (Por recurrencia histórica del cliente)
            if c_up:
                if c_up in history_client_resp:
                    hist_ana = history_client_resp[c_up]
                    if hist_ana and hist_ana != 'Sin Asignar':
                        metodo = f"Historial Cliente ({hist_ana})"
                        if author_sap and author_sap.lower() != hist_ana.lower():
                            metodo += f" [Apoyo facturación SAP: {author_sap}]"
                        return hist_ana, author_desc, metodo
                for k, v in history_client_resp.items():
                    if k in c_up or c_up in k:
                        if v and v != 'Sin Asignar':
                            metodo = f"Historial Cliente ({v})"
                            if author_sap and author_sap.lower() != v.lower():
                                metodo += f" [Apoyo facturación SAP: {author_sap}]"
                            return v, author_desc, metodo

            # Regla 3: Si no hay cartera ni historial, asignar al autor SAP si es de Congelado
            if author_code in congelado_authors_set:
                return author_sap, author_sap, f"Facturado por Congelado ({author_sap})"

            return "Sin Asignar", author_desc, "Por definir"

        seen_flow_peds = set()
        for r in range(2, ws_fl.max_row + 1):
            f_desp = ws_fl.cell(row=r, column=1).value
            f_fac = ws_fl.cell(row=r, column=2).value
            ped = ws_fl.cell(row=r, column=4).value
            entrega = ws_fl.cell(row=r, column=5).value
            cli = clean_client_name(ws_fl.cell(row=r, column=6).value)
            folio = normalize_sap_doc(ws_fl.cell(row=r, column=7).value)
            doc_fac = normalize_sap_doc(ws_fl.cell(row=r, column=8).value)
            desc = str(ws_fl.cell(row=r, column=10).value or '').strip()
            cajas = ws_fl.cell(row=r, column=11).value
            kilos = ws_fl.cell(row=r, column=12).value
            
            factura_flujo = doc_fac or folio
            if ped:
                ped_str = str(ped).strip().lstrip('0')
                if ped_str not in active_peds_set and ped_str not in seen_flow_peds:
                    seen_flow_peds.add(ped_str)
                    analyst, author_sap, method = resolve_flow_analyst(cli, factura_flujo, ped_str=ped_str)
                    
                    # FILTRADO FRESCO: Si fue tramitado por Fresco, se elimina del Control Congelado
                    if analyst == 'FRESCO_OMITIR' or author_sap in fresco_authors_set:
                        flujo_fresco_omitidos += 1
                        continue

                    # REGLA DE CORTE HISTÓRICO: Excluir zarpes anteriores a Junio 2026
                    f_dt = parse_comex_date(f_desp)
                    if f_dt and f_dt < datetime.date(2026, 6, 1):
                        continue

                    # Evaluación precisa de hitos operativos (48h BL naviera, 4d hábiles draft, 15d finales)
                    m = evaluate_order_milestones(f_desp, None, None, factura_flujo, None, today_dt=today_dt, holidays=feriados_set, cli=cli)

                    flow_orders.append({
                        "id": 10000 + len(flow_orders) + 1,
                        "tipo": "Flujo / Por Zarpar",
                        "estatus": m["estatus"],
                        "pedido": ped_str,
                        "codigo_cliente": find_sap_code_for_client(cli),
                        "responsable": analyst,
                        "autor_sap": author_sap,
                        "metodo_asignacion": method,
                        "cliente": cli,
                        "categoria_cliente": m["categoria_cliente"],
                        "formato_liberacion": m["formato_liberacion"],
                        "nave": "Flujo SAP (Por definir)",
                        "puerto": "",
                        "zarpe": m["zarpe"],
                        "fecha_factura": factura_date.get(doc_fac or folio, '') or (str(f_fac)[:10] if f_fac else ''),
                        "doc_factura": doc_fac,
                        "factura": factura_flujo,
                        "descripcion": desc,
                        "cajas": str(cajas).strip() if cajas else '0',
                        "kilos": str(kilos).strip() if kilos else '0',
                        "bl_esperado": m["bl_esperado"],
                        "estado_bl": m["estado_bl"],
                        "alerta_bl": m["alerta_bl"],
                        "dias_atraso_bl": m["dias_atraso_bl"],
                        "draft_real": "",
                        "draft_esperado": m["draft_esperado"],
                        "draft_plazo": m["draft_esperado"] or "",
                        "finales_real": "",
                        "finales_esperado": m["finales_esperado"],
                        "estado_draft": m["estado_draft"],
                        "estado_finales": m["estado_finales"],
                        "dias_atraso_draft": m["dias_atraso_draft"],
                        "dias_atraso_finales": m["dias_atraso_finales"],
                        "nro_dhl": m["nro_dhl"],
                        "fecha_dhl": m["fecha_dhl"],
                        "condicion_venta": "",
                        "swb": "",
                        "swb_liberacion": "",
                        "release": "",
                        "estado_flujo": "Facturado en SAP" if doc_fac else "Creado en SAP (Por Zarpar/Facturar)"
                    })

        print(f"   [OK] {len(flow_orders)} pedidos de flujo Congelado integrados ({flujo_fresco_omitidos} pedidos de FRESCO eliminados).")

    # 6. Triangulación con Vista SQL de Azure Synapse (pedidos_sql.xlsx / pedidos_sql.csv)
    sql_file = "pedidos_sql.xlsx" if os.path.exists("pedidos_sql.xlsx") else ("pedidos_sql.csv" if os.path.exists("pedidos_sql.csv") else None)
    if sql_file:
        print(f"-> Triangulando con datos en vivo de SQL Synapse ({sql_file})...")
        try:
            if sql_file.endswith('.csv'):
                df_sql = pd.read_csv(sql_file)
            else:
                df_sql = pd.read_excel(sql_file)

            active_dict = {s['pedido']: s for s in active_shipments}
            flow_dict = {s['pedido']: s for s in flow_orders}
            sql_nuevos_count = 0
            sql_actualizados_count = 0
            sql_fresco_omitidos = 0

            for _, row in df_sql.iterrows():
                ped_val = str(row.get('Vbeln') or row.get('pedido') or '').strip().lstrip('0')
                if not ped_val:
                    continue

                cli_val = clean_client_name(row.get('Nombre Cliente') or row.get('Name1') or '')
                cod_cli_val = str(row.get('Codigo Cliente') or row.get('Kunnr') or '').strip().lstrip('0')
                if cod_cli_val.lower() in ('nan', 'none', '-'): cod_cli_val = ''
                dest_val = str(row.get('Destinatario') or row.get('NomDmercaderia') or '').strip()
                if dest_val.lower() in ('nan', 'none', '-'): dest_val = ''
                
                contrato_val = str(row.get('Nro Contrato') or row.get('Bstkd') or '').strip()
                if contrato_val.lower() in ('nan', 'none', '-', '0'): contrato_val = ''
                
                sociedad_val = str(row.get('Sociedad') or row.get('Vkorg') or '').strip()
                if 'S170' in sociedad_val: sociedad_val = 'S170'
                elif 'S070' in sociedad_val: sociedad_val = 'S070'
                elif sociedad_val.lower() in ('nan', 'none', '-'): sociedad_val = ''

                cond_venta_val = str(row.get('Condicion Venta') or row.get('Zlschtxt') or '').strip()
                if cond_venta_val.lower() in ('nan', 'none', '-', '0'): cond_venta_val = ''
                if cond_venta_val.lower() == 'intercompany': cond_venta_val = 'Crédito'

                inco_val = str(row.get('Incoterm') or row.get('Inco1') or '').strip()
                if inco_val.lower() in ('nan', 'none', '-'): inco_val = ''

                nave_val = str(row.get('Nave') or '').strip()
                naviera_val = str(row.get('Naviera') or row.get('TxtTransportista') or '').strip()
                
                # Folio SAP (match exacto con transacción VF05N: 974747xxx) y Folio SII (DTE Legal: 177xxx)
                folio_sap = normalize_sap_doc(row.get('Folio SAP') or row.get('Fkvbeln'))
                if folio_sap.lower() in ('no facturado', '0'): folio_sap = ''
                factura_sii = normalize_sap_doc(row.get('Nro Factura') or row.get('Fkfolio'))
                if factura_sii.lower() in ('no facturado', '0'): factura_sii = ''
                factura_para_hitos = factura_sii or folio_sap

                etd_dt = parse_comex_date(row.get('ETD') or row.get('etd'))
                etd_val = etd_dt.strftime('%Y-%m-%d') if etd_dt else ''
                eta_dt = parse_comex_date(row.get('ETA') or row.get('etaPredictionP44'))
                eta_val = eta_dt.strftime('%Y-%m-%d') if eta_dt else ''
                cont_val = str(row.get('Nro Contenedor') or row.get('NumCont') or '').strip()
                if cont_val.lower() in ('nan', 'none', '-'): cont_val = ''
                bl_val = str(row.get('Nro BL') or row.get('Bl') or '').strip()
                if bl_val.lower() in ('nan', 'none', '-'): bl_val = ''
                puerto_des_val = str(row.get('Puerto Destino') or row.get('PuertoDes') or '').strip()
                if puerto_des_val.lower() in ('nan', 'none', '-'): puerto_des_val = ''

                # Si ya existe en activos, enriquecemos con Synapse en vivo y Project44
                if ped_val in active_dict:
                    item = active_dict[ped_val]
                    if etd_val and not item.get('zarpe'): item['zarpe'] = etd_val
                    if eta_val and not item.get('eta'): item['eta'] = eta_val
                    if nave_val and (not item.get('nave') or item['nave'] in ('Flujo SAP (Por definir)', '')):
                        item['nave'] = nave_val
                    if not item.get('puerto') and nave_val: item['puerto'] = nave_to_port.get(nave_val.upper(), '')
                    if cont_val and (not item.get('contenedor') or item['contenedor'] in ('nan', 'None', '-', '')):
                        item['contenedor'] = cont_val
                    if bl_val and (not item.get('bl') or item['bl'] in ('nan', 'None', '-', '')):
                        item['bl'] = bl_val
                    if naviera_val and (not item.get('naviera') or item['naviera'] in ('nan', 'None', '-', '')):
                        item['naviera'] = naviera_val
                    if contrato_val and (not item.get('contrato') or item['contrato'] in ('nan', 'None', '-', '')):
                        item['contrato'] = contrato_val
                    if sociedad_val and (not item.get('sociedad') or item['sociedad'] in ('nan', 'None', '-', '')):
                        item['sociedad'] = sociedad_val
                    if cond_venta_val and (not item.get('condicion_venta') or str(item.get('condicion_venta')).lower() in ('intercompany', 'nan', 'none', '-')):
                        item['condicion_venta'] = cond_venta_val
                    if inco_val and (not item.get('incoterm') or item['incoterm'] in ('nan', 'None', '-', '')):
                        item['incoterm'] = inco_val
                    if cod_cli_val and (not item.get('codigo_cliente') or item.get('codigo_cliente') in ('Desconocido', 'nan', '')):
                        item['codigo_cliente'] = cod_cli_val
                    if dest_val and (not item.get('destinatario') or item['destinatario'] in ('nan', 'None', '-', '')):
                        item['destinatario'] = dest_val
                    if puerto_des_val and (not item.get('puerto_destino') or item['puerto_destino'] in ('nan', 'None', '-', '')):
                        item['puerto_destino'] = puerto_des_val
                    if folio_sap: item['doc_factura'] = folio_sap
                    if factura_sii:
                        item['factura'] = factura_sii
                    elif item.get('factura'):
                        f_cur = str(item['factura']).strip()
                        if f_cur.startswith('9') and len(f_cur) >= 8:
                            if not item.get('doc_factura'): item['doc_factura'] = f_cur
                            item['factura'] = ''
                    
                    # Match con VF05N para autor y fecha
                    fact_key = normalize_sap_doc(item.get('doc_factura') or folio_sap)
                    if fact_key:
                        aut_code = factura_author.get(fact_key, '')
                        if aut_code:
                            aut_name = EQUIPOS_AUTORES.get(aut_code, {}).get('analista', aut_code)
                            item['autor_sap'] = aut_name
                        if fact_key in factura_date and not item.get('fecha_factura'):
                            item['fecha_factura'] = factura_date[fact_key]
                            
                    item['origen_sql'] = 'Enriquecido con Synapse/P44'
                    # Re-evaluar hitos operativos con los nuevos datos de zarpe, BL o condición de venta
                    factura_eval = item.get('factura') or item.get('doc_factura') or factura_para_hitos
                    m = evaluate_order_milestones(item.get('zarpe'), item.get('draft_real'), item.get('finales_real'), factura_eval, item.get('bl'), dhl=item.get('fecha_dhl') or item.get('dhl'), swb=item.get('swb'), fin_e=item.get('finales_esperado'), today_dt=today_dt, holidays=feriados_set, condicion_venta=item.get('condicion_venta'), nro_dhl_explicit=item.get('nro_dhl'), cli=item.get('cliente'), co_aprobado=item.get('co_aprobado'), formato_liberacion=item.get('formato_liberacion'))
                    item['estatus'] = m['estatus']
                    item['categoria_cliente'] = m['categoria_cliente']
                    item['formato_liberacion'] = m['formato_liberacion']
                    if m.get('swb_liberacion'):
                        item['swb_liberacion'] = m['swb_liberacion']
                        item['swb'] = m['swb_liberacion']
                        item['release'] = m['swb_liberacion']
                    item['bl_esperado'] = m['bl_esperado']
                    item['estado_bl'] = m['estado_bl']
                    item['alerta_bl'] = m['alerta_bl']
                    item['dias_atraso_bl'] = m['dias_atraso_bl']
                    item['draft_esperado'] = m['draft_esperado']
                    item['draft_plazo'] = m['draft_esperado'] or item.get('draft_plazo', '')
                    item['estado_draft'] = m['estado_draft']
                    item['dias_atraso_draft'] = m['dias_atraso_draft']
                    item['finales_esperado'] = m['finales_esperado']
                    item['estado_finales'] = m['estado_finales']
                    item['dias_atraso_finales'] = m['dias_atraso_finales']
                    item['nro_dhl'] = m['nro_dhl']
                    item['fecha_dhl'] = m['fecha_dhl']
                    sql_actualizados_count += 1
                elif ped_val in flow_dict:
                    item = flow_dict[ped_val]
                    if etd_val and not item.get('zarpe'): item['zarpe'] = etd_val
                    if eta_val and not item.get('eta'): item['eta'] = eta_val
                    if nave_val and (not item.get('nave') or item['nave'] in ('Flujo SAP (Por definir)', '')):
                        item['nave'] = nave_val
                    if not item.get('puerto') and nave_val: item['puerto'] = nave_to_port.get(nave_val.upper(), '')
                    if cont_val and (not item.get('contenedor') or item['contenedor'] in ('nan', 'None', '-', '')):
                        item['contenedor'] = cont_val
                    if bl_val and (not item.get('bl') or item['bl'] in ('nan', 'None', '-', '')):
                        item['bl'] = bl_val
                    if naviera_val and (not item.get('naviera') or item['naviera'] in ('nan', 'None', '-', '')):
                        item['naviera'] = naviera_val
                    if contrato_val and (not item.get('contrato') or item['contrato'] in ('nan', 'None', '-', '')):
                        item['contrato'] = contrato_val
                    if sociedad_val and (not item.get('sociedad') or item['sociedad'] in ('nan', 'None', '-', '')):
                        item['sociedad'] = sociedad_val
                    if cond_venta_val and (not item.get('condicion_venta') or str(item.get('condicion_venta')).lower() in ('intercompany', 'nan', 'none', '-')):
                        item['condicion_venta'] = cond_venta_val
                    if inco_val and (not item.get('incoterm') or item['incoterm'] in ('nan', 'None', '-', '')):
                        item['incoterm'] = inco_val
                    if cod_cli_val and (not item.get('codigo_cliente') or item.get('codigo_cliente') in ('Desconocido', 'nan', '')):
                        item['codigo_cliente'] = cod_cli_val
                    if dest_val and (not item.get('destinatario') or item['destinatario'] in ('nan', 'None', '-', '')):
                        item['destinatario'] = dest_val
                    if puerto_des_val and (not item.get('puerto_destino') or item['puerto_destino'] in ('nan', 'None', '-', '')):
                        item['puerto_destino'] = puerto_des_val
                    if (not item.get('cliente') or item['cliente'] in ('nan', 'None', '')) and cli_val:
                        item['cliente'] = cli_val
                    if folio_sap: item['doc_factura'] = folio_sap
                    if factura_sii:
                        item['factura'] = factura_sii
                    elif item.get('factura'):
                        f_cur = str(item['factura']).strip()
                        if f_cur.startswith('9') and len(f_cur) >= 8:
                            if not item.get('doc_factura'): item['doc_factura'] = f_cur
                            item['factura'] = ''
                    
                    # Match con VF05N para autor y fecha
                    fact_key = normalize_sap_doc(item.get('doc_factura') or folio_sap)
                    if fact_key:
                        aut_code = factura_author.get(fact_key, '')
                        if aut_code:
                            aut_name = EQUIPOS_AUTORES.get(aut_code, {}).get('analista', aut_code)
                            item['autor_sap'] = aut_name
                        if fact_key in factura_date and not item.get('fecha_factura'):
                            item['fecha_factura'] = factura_date[fact_key]
                            
                    # Si no tenía analista asignado, re-evaluar con la factura
                    if not item.get('responsable') or item.get('responsable') == 'Sin Asignar':
                        ana_res, aut_res, met_res = resolve_flow_analyst(item.get('cliente'), fact_key, ped_str=ped_val)
                        if ana_res and ana_res != 'Sin Asignar':
                            item['responsable'] = ana_res
                            item['metodo_asignacion'] = met_res
                            
                    item['origen_sql'] = 'Enriquecido con Synapse/P44'
                    # Re-evaluar hitos operativos con los nuevos datos de zarpe, BL o condición de venta
                    factura_eval = item.get('factura') or item.get('doc_factura') or factura_para_hitos
                    m = evaluate_order_milestones(item.get('zarpe'), item.get('draft_real'), item.get('finales_real'), factura_eval, item.get('bl'), dhl=item.get('fecha_dhl') or item.get('dhl'), swb=item.get('swb'), fin_e=item.get('finales_esperado'), today_dt=today_dt, holidays=feriados_set, condicion_venta=item.get('condicion_venta'), nro_dhl_explicit=item.get('nro_dhl'), cli=item.get('cliente'), co_aprobado=item.get('co_aprobado'), formato_liberacion=item.get('formato_liberacion'))
                    item['estatus'] = m['estatus']
                    item['categoria_cliente'] = m['categoria_cliente']
                    item['formato_liberacion'] = m['formato_liberacion']
                    if m.get('swb_liberacion'):
                        item['swb_liberacion'] = m['swb_liberacion']
                        item['swb'] = m['swb_liberacion']
                        item['release'] = m['swb_liberacion']
                    item['bl_esperado'] = m['bl_esperado']
                    item['estado_bl'] = m['estado_bl']
                    item['alerta_bl'] = m['alerta_bl']
                    item['dias_atraso_bl'] = m['dias_atraso_bl']
                    item['draft_esperado'] = m['draft_esperado']
                    item['draft_plazo'] = m['draft_esperado'] or item.get('draft_plazo', '')
                    item['estado_draft'] = m['estado_draft']
                    item['dias_atraso_draft'] = m['dias_atraso_draft']
                    item['finales_esperado'] = m['finales_esperado']
                    item['estado_finales'] = m['estado_finales']
                    item['dias_atraso_finales'] = m['dias_atraso_finales']
                    item['nro_dhl'] = m['nro_dhl']
                    item['fecha_dhl'] = m['fecha_dhl']
                    sql_actualizados_count += 1
                else:
                    # Filtro SQL: Solo ingresar pedidos relevantes de la temporada oficial (Junio 2026 en adelante)
                    p_num = None
                    try:
                        p_num = int(ped_val)
                        if p_num < 40566606:
                            continue # Excluir estrictamente histórico de años o temporadas anteriores
                    except Exception:
                        pass

                    analyst, author_sap, method = resolve_flow_analyst(cli_val, folio_sap or factura_sii, ped_str=ped_val, cli_code=cod_cli_val)
                    if analyst == 'FRESCO_OMITIR' or author_sap in fresco_authors_set:
                        sql_fresco_omitidos += 1
                        continue

                    # REGLA DE CORTE HISTÓRICO: Excluir zarpes anteriores a Junio 2026
                    etd_dt = parse_comex_date(etd_val)
                    if etd_dt and etd_dt < datetime.date(2026, 6, 1):
                        continue
                    # Si el pedido es anterior a Junio 2026 (< 40586946) y no tiene zarpe confirmado, descartar por reserva caduca/cancelada
                    if p_num is not None and p_num < 40586946 and not etd_val:
                        continue

                    m = evaluate_order_milestones(etd_val, None, None, factura_para_hitos, bl_val, today_dt=today_dt, holidays=feriados_set, condicion_venta=cond_venta_val, cli=cli_val)
                    
                    fact_key = normalize_sap_doc(folio_sap)
                    f_date = factura_date.get(fact_key, '') if fact_key else ''

                    new_order = {
                        "id": 20000 + len(flow_orders) + 1,
                        "tipo": "Flujo / Por Zarpar" if m["estatus"] == "Pendiente de facturar" else "Flujo Futuro",
                        "estatus": m["estatus"],
                        "pedido": ped_val,
                        "codigo_cliente": cod_cli_val or find_sap_code_for_client(cli_val),
                        "responsable": analyst,
                        "autor_sap": author_sap,
                        "metodo_asignacion": f"SQL Synapse + {method}",
                        "cliente": cli_val,
                        "destinatario": dest_val,
                        "sociedad": sociedad_val,
                        "contrato": contrato_val,
                        "condicion_venta": cond_venta_val,
                        "incoterm": inco_val,
                        "nave": nave_val or "Flujo SAP (Por definir)",
                        "puerto": nave_to_port.get(nave_val.upper(), '') if nave_val else '',
                        "puerto_destino": puerto_des_val,
                        "naviera": naviera_val,
                        "contenedor": cont_val,
                        "bl": bl_val,
                        "zarpe": m["zarpe"],
                        "eta": eta_val,
                        "fecha_factura": f_date,
                        "doc_factura": folio_sap,
                        "factura": factura_sii,
                        "descripcion": "Ingesta automática desde Azure Synapse",
                        "cajas": "0",
                        "kilos": "0",
                        "bl_esperado": m["bl_esperado"],
                        "estado_bl": m["estado_bl"],
                        "alerta_bl": m["alerta_bl"],
                        "dias_atraso_bl": m["dias_atraso_bl"],
                        "draft_real": "",
                        "draft_esperado": m["draft_esperado"],
                        "draft_plazo": m["draft_esperado"] or "",
                        "certificados": "",
                        "co_aprobado": "",
                        "finales_real": "",
                        "finales_esperado": m["finales_esperado"],
                        "estado_draft": m["estado_draft"],
                        "estado_finales": m["estado_finales"],
                        "dias_atraso_draft": m["dias_atraso_draft"],
                        "dias_atraso_finales": m["dias_atraso_finales"],
                        "categoria_cliente": m.get("categoria_cliente", "Comercial"),
                        "formato_liberacion": m.get("formato_liberacion", "OBL"),
                        "swb": "",
                        "swb_liberacion": "",
                        "nro_dhl": m["nro_dhl"],
                        "fecha_dhl": m["fecha_dhl"],
                        "estado_flujo": "Creado en SAP (SQL Synapse en vivo)",
                        "origen_sql": "Nuevo desde Azure Synapse"
                    }
                    flow_orders.append(new_order)
                    flow_dict[ped_val] = new_order
                    sql_nuevos_count += 1

            print(f"   [OK] SQL Triangulado: {sql_actualizados_count} pedidos enriquecidos, {sql_nuevos_count} pedidos nuevos ingresados ({sql_fresco_omitidos} pedidos Fresco omitidos).")
        except Exception as e_sql:
            print(f"   [AVISO] Error al triangular SQL: {e_sql}")

    # 6.2. Enriquecimiento con Planilla Maestra de Exportaciones (Exportaciones gestion Pedidos de Vta.xlsx)
    exp_file = "Exportaciones gestion Pedidos de Vta.xlsx"
    if os.path.exists(exp_file):
        print(f"-> Enriqueciendo con Planilla Maestra de Exportaciones ({exp_file})...")
        try:
            mapeo_pago_exp = {
                'CTA CTE O CRED.DIRECTO': 'Crédito',
                'PAGO ANTIC. – PAGO C/COPIA DOC': 'Anticipo',
                'PAGO ANTIC. - PAGO C/COPIA DOC': 'Anticipo',
                'PAGO ANTICIPADO': 'Anticipo',
                'PAGO C/COPIA DOCTO.': 'PAGO C/COPIA DOCTO.',
                'TRANSFERENCIA': 'Transferencia',
                'CARTA CREDITO': 'CARTA CREDITO'
            }
            mapeo_sociedad_exp = {
                'S170 AQUACHILE S.A.': 'S170',
                'S070 PESQUERA LOS FIORDOS': 'S070'
            }

            excel_exp_obj = pd.ExcelFile(exp_file)
            if 'Hoja1' in excel_exp_obj.sheet_names:
                df_h1 = excel_exp_obj.parse('Hoja1', header=None)
                for _, row in df_h1.iterrows():
                    k_val = row[0] if len(row) > 0 else None
                    v_val = row[1] if len(row) > 1 else None
                    if pd.notna(k_val) and pd.notna(v_val):
                        k_s = str(k_val).strip()
                        v_s = str(v_val).strip()
                        if 'Vía de Pago' not in k_s and 'Organización' not in k_s and 'Total' not in k_s and '(en blanco)' not in k_s:
                            if k_s.startswith('S') and ' ' in k_s:
                                mapeo_sociedad_exp[k_s] = v_s
                            else:
                                mapeo_pago_exp[k_s] = v_s

            df_exp = excel_exp_obj.parse('Data')
            ped_exp_dict = {}

            for row in df_exp.itertuples(index=False):
                p_raw = getattr(row, '_0', getattr(row, 'Nro__Pedido', '')) if len(row) > 0 else ''
                ped = str(p_raw or '').strip().lstrip('0')
                if not ped or ped == 'nan':
                    continue

                if ped not in ped_exp_dict:
                    ped_exp_dict[ped] = {
                        'bl': '', 'naviera': '', 'sociedad': '', 'contrato': '',
                        'condicion_venta': '', 'incoterm': '', 'nave': '',
                        'contenedores': [], 'cliente': '', 'codigo_cliente': '',
                        'factura': '', 'doc_factura': ''
                    }
                p_entry = ped_exp_dict[ped]

                if len(row) > 7 and pd.notna(row[7]):
                    bl_v = str(row[7]).strip()
                    if bl_v not in ('nan', 'None', '', '-', '0') and not p_entry['bl']:
                        p_entry['bl'] = bl_v

                if len(row) > 14 and pd.notna(row[14]):
                    nav_v = str(row[14]).strip()
                    if nav_v not in ('nan', 'None', '', '-') and not p_entry['naviera']:
                        p_entry['naviera'] = nav_v

                if len(row) > 5 and pd.notna(row[5]):
                    org_v = str(row[5]).strip()
                    if org_v not in ('nan', 'None', '', '-') and not p_entry['sociedad']:
                        p_entry['sociedad'] = mapeo_sociedad_exp.get(org_v, 'S170' if 'S170' in org_v else ('S070' if 'S070' in org_v else org_v[:4]))

                if len(row) > 8 and pd.notna(row[8]):
                    cont_v = str(row[8]).strip()
                    if cont_v not in ('nan', 'None', '', '-') and not p_entry['contrato']:
                        p_entry['contrato'] = cont_v

                if len(row) > 4 and pd.notna(row[4]):
                    pago_v = str(row[4]).strip()
                    if pago_v not in ('nan', 'None', '', '-') and not p_entry['condicion_venta']:
                        p_up = pago_v.upper()
                        if 'CRED' in p_up or 'CTA CTE' in p_up: p_entry['condicion_venta'] = 'Crédito'
                        elif 'ANTIC' in p_up: p_entry['condicion_venta'] = 'Anticipo'
                        elif 'TRANSF' in p_up: p_entry['condicion_venta'] = 'Transferencia'
                        elif 'CARTA' in p_up: p_entry['condicion_venta'] = 'CARTA CREDITO'
                        else: p_entry['condicion_venta'] = mapeo_pago_exp.get(pago_v, pago_v)

                if len(row) > 11 and pd.notna(row[11]):
                    inc_v = str(row[11]).strip()
                    if inc_v not in ('nan', 'None', '', '-') and not p_entry['incoterm']:
                        p_entry['incoterm'] = inc_v

                if len(row) > 12 and pd.notna(row[12]):
                    nave_v = str(row[12]).strip()
                    if nave_v not in ('nan', 'None', '', '-') and not p_entry['nave']:
                        p_entry['nave'] = nave_v

                if len(row) > 13 and pd.notna(row[13]):
                    cnt_v = str(row[13]).strip()
                    if cnt_v not in ('nan', 'None', '', '-') and cnt_v not in p_entry['contenedores']:
                        p_entry['contenedores'].append(cnt_v)

                if len(row) > 2 and pd.notna(row[2]):
                    cli_v = str(row[2]).strip()
                    if cli_v not in ('nan', 'None', '', '-') and not p_entry['cliente']:
                        p_entry['cliente'] = cli_v
                if len(row) > 1 and pd.notna(row[1]):
                    cod_v = str(row[1]).strip()
                    if cod_v not in ('nan', 'None', '', '-') and not p_entry['codigo_cliente']:
                        p_entry['codigo_cliente'] = cod_v

                if len(row) > 9 and pd.notna(row[9]):
                    fac_sii_v = normalize_sap_doc(row[9])
                    if fac_sii_v and fac_sii_v not in ('0', '0.0', 'No facturado'):
                        p_entry['factura'] = fac_sii_v
                if len(row) > 10 and pd.notna(row[10]):
                    fac_sap_v = normalize_sap_doc(row[10])
                    if fac_sap_v and fac_sap_v not in ('0', '0.0', 'No facturado'):
                        p_entry['doc_factura'] = fac_sap_v

            for p_dict in ped_exp_dict.values():
                p_dict['contenedor'] = ', '.join(p_dict['contenedores'])
                del p_dict['contenedores']

            exp_activos_count = 0
            exp_flujo_count = 0

            for item in active_shipments:
                p_k = str(item.get('pedido', '')).strip().lstrip('0')
                if p_k in ped_exp_dict:
                    info = ped_exp_dict[p_k]
                    if info['sociedad']: item['sociedad'] = info['sociedad']
                    if info['contrato']: item['contrato'] = info['contrato']
                    if info['condicion_venta']: item['condicion_venta'] = info['condicion_venta']
                    if info['incoterm']: item['incoterm'] = info['incoterm']
                    if info['naviera'] and (not item.get('naviera') or item['naviera'] in ('nan', 'None', '-', '')):
                        item['naviera'] = info['naviera']
                    if info['bl'] and (not item.get('bl') or item['bl'] in ('nan', 'None', '-', '')):
                        item['bl'] = info['bl']
                    if info['contenedor'] and (not item.get('contenedor') or item['contenedor'] in ('nan', 'None', '-', '')):
                        item['contenedor'] = info['contenedor']
                    if info['nave'] and (not item.get('nave') or item['nave'] in ('Flujo SAP (Por definir)', '')):
                        item['nave'] = info['nave']

                    # Enriquecimiento de Folio SII y Folio SAP
                    if info['doc_factura']:
                        item['doc_factura'] = info['doc_factura']
                    if info['factura']:
                        item['factura'] = info['factura']
                    elif item.get('factura'):
                        f_cur = str(item['factura']).strip()
                        if f_cur.startswith('9') and len(f_cur) >= 8:
                            if not item.get('doc_factura'): item['doc_factura'] = f_cur
                            item['factura'] = ''

                    # Verificar si con el doc_factura podemos obtener autor SAP o fecha
                    f_sap_key = normalize_sap_doc(item.get('doc_factura'))
                    if f_sap_key:
                        if not item.get('autor_sap') or item.get('autor_sap') == 'Sin Factura':
                            aut_c = factura_author.get(f_sap_key, '')
                            if aut_c:
                                item['autor_sap'] = EQUIPOS_AUTORES.get(aut_c, {}).get('analista', aut_c)
                        if f_sap_key in factura_date and not item.get('fecha_factura'):
                            item['fecha_factura'] = factura_date[f_sap_key]

                    # Si el pedido estaba "Pendiente de facturar" y ahora tiene factura, re-evaluar
                    if item.get('estatus') in ('Pendiente de facturar', 'Por Despachar') and (item.get('factura') or item.get('doc_factura')):
                        m = evaluate_order_milestones(item.get('zarpe'), item.get('draft_real'), item.get('finales_real'), item.get('factura') or item.get('doc_factura'), item.get('bl'), dhl=item.get('fecha_dhl') or item.get('dhl'), swb=item.get('swb'), fin_e=item.get('finales_esperado'), today_dt=today_dt, holidays=feriados_set, condicion_venta=item.get('condicion_venta'), nro_dhl_explicit=item.get('nro_dhl'), cli=item.get('cliente'), co_aprobado=item.get('co_aprobado'), formato_liberacion=item.get('formato_liberacion'))
                        item['estatus'] = m['estatus']
                        item['draft_plazo'] = m['draft_esperado'] or item.get('draft_plazo', '')
                        item['draft_esperado'] = m['draft_esperado'] or item.get('draft_esperado', '')
                        item['estado_draft'] = m['estado_draft']
                        item['dias_atraso_draft'] = m['dias_atraso_draft']
                        item['finales_esperado'] = m['finales_esperado']
                        item['estado_finales'] = m['estado_finales']
                        item['dias_atraso_finales'] = m['dias_atraso_finales']
                        if m.get('nro_dhl'): item['nro_dhl'] = m['nro_dhl']
                        if m.get('fecha_dhl'): item['fecha_dhl'] = m['fecha_dhl']

                    exp_activos_count += 1

            for item in flow_orders:
                p_k = str(item.get('pedido', '')).strip().lstrip('0')
                if p_k in ped_exp_dict:
                    info = ped_exp_dict[p_k]
                    if info['sociedad']: item['sociedad'] = info['sociedad']
                    if info['contrato']: item['contrato'] = info['contrato']
                    if info['condicion_venta']: item['condicion_venta'] = info['condicion_venta']
                    if info['incoterm']: item['incoterm'] = info['incoterm']
                    if info['naviera'] and (not item.get('naviera') or item['naviera'] in ('nan', 'None', '-', '')):
                        item['naviera'] = info['naviera']
                    if info['bl'] and (not item.get('bl') or item['bl'] in ('nan', 'None', '-', '')):
                        item['bl'] = info['bl']
                    if info['contenedor'] and (not item.get('contenedor') or item['contenedor'] in ('nan', 'None', '-', '')):
                        item['contenedor'] = info['contenedor']
                    if info['nave'] and (not item.get('nave') or item['nave'] in ('Flujo SAP (Por definir)', '')):
                        item['nave'] = info['nave']
                        if not item.get('puerto'):
                            item['puerto'] = nave_to_port.get(info['nave'].upper(), '')

                    # Enriquecimiento de Folio SII y Folio SAP
                    if info['doc_factura']:
                        item['doc_factura'] = info['doc_factura']
                    if info['factura']:
                        item['factura'] = info['factura']
                    elif item.get('factura'):
                        f_cur = str(item['factura']).strip()
                        if f_cur.startswith('9') and len(f_cur) >= 8:
                            if not item.get('doc_factura'): item['doc_factura'] = f_cur
                            item['factura'] = ''

                    f_sap_key = normalize_sap_doc(item.get('doc_factura'))
                    if f_sap_key:
                        if not item.get('autor_sap') or item.get('autor_sap') == 'Sin Factura':
                            aut_c = factura_author.get(f_sap_key, '')
                            if aut_c:
                                item['autor_sap'] = EQUIPOS_AUTORES.get(aut_c, {}).get('analista', aut_c)
                        if f_sap_key in factura_date and not item.get('fecha_factura'):
                            item['fecha_factura'] = factura_date[f_sap_key]

                    if (not item.get('cliente') or item['cliente'] in ('nan', 'None')) and info['cliente']:
                        item['cliente'] = info['cliente']
                        item['codigo_cliente'] = info['codigo_cliente']
                        if item.get('responsable') == 'Sin Asignar':
                            ana_res, aut_res, met_res = resolve_flow_analyst(info['cliente'], item.get('doc_factura'), ped_str=p_k)
                            if ana_res and ana_res != 'Sin Asignar':
                                item['responsable'] = ana_res
                                item['metodo_asignacion'] = met_res

                    if item.get('estatus') in ('Pendiente de facturar', 'Por Despachar') and (item.get('factura') or item.get('doc_factura')):
                        m = evaluate_order_milestones(item.get('zarpe'), item.get('draft_real'), item.get('finales_real'), item.get('factura') or item.get('doc_factura'), item.get('bl'), dhl=item.get('fecha_dhl') or item.get('dhl'), swb=item.get('swb'), fin_e=item.get('finales_esperado'), today_dt=today_dt, holidays=feriados_set, condicion_venta=item.get('condicion_venta'), nro_dhl_explicit=item.get('nro_dhl'), cli=item.get('cliente'), co_aprobado=item.get('co_aprobado'), formato_liberacion=item.get('formato_liberacion'))
                        item['estatus'] = m['estatus']
                        item['draft_plazo'] = m['draft_esperado'] or item.get('draft_plazo', '')
                        item['draft_esperado'] = m['draft_esperado'] or item.get('draft_esperado', '')
                        item['estado_draft'] = m['estado_draft']
                        item['dias_atraso_draft'] = m['dias_atraso_draft']
                        item['finales_esperado'] = m['finales_esperado']
                        item['estado_finales'] = m['estado_finales']
                        item['dias_atraso_finales'] = m['dias_atraso_finales']
                        if m.get('nro_dhl'): item['nro_dhl'] = m['nro_dhl']
                        if m.get('fecha_dhl'): item['fecha_dhl'] = m['fecha_dhl']

                    exp_flujo_count += 1

            print(f"   [OK] Planilla Maestra Exportaciones integrada: {exp_activos_count} activos y {exp_flujo_count} pedidos de flujo enriquecidos (Naviera, BL, Sociedad, Contrato, Condición Venta e Incoterm).")
        except Exception as e_exp:
            print(f"   [AVISO] Error al procesar {exp_file}: {e_exp}")

    # 6.5. Integrar Modificaciones y Pedidos Manuales si existe cambios_manuales.json
    json_manual_path = "cambios_manuales.json"
    if os.path.exists(json_manual_path):
        try:
            with open(json_manual_path, "r", encoding="utf-8") as f_man:
                man_data = json.load(f_man)

            # Integrar nuevos pedidos manuales
            man_new = man_data.get("newOrders", [])
            new_added_count = 0
            for no in man_new:
                p_str = str(no.get("pedido", "")).strip()
                if p_str and not any(str(x.get("pedido", "")).strip() == p_str for x in active_shipments + flow_orders):
                    if no.get("estatus") in ("Pendiente de facturar", "Por Despachar"):
                        flow_orders.insert(0, no)
                    else:
                        active_shipments.insert(0, no)
                    new_added_count += 1

            # Integrar modificaciones de campos
            man_edits = man_data.get("edits", {})
            edits_applied_count = 0
            for p_str, ed in man_edits.items():
                target = next((x for x in active_shipments + flow_orders if str(x.get("pedido", "")).strip() == str(p_str).strip()), None)
                if target:
                    for k, v in ed.items():
                        if k not in ("pedido", "timestamp"):
                            target[k] = v
                    edits_applied_count += 1

            print(f"-> Integrando cambios guardados manualmente ({json_manual_path})...")
            print(f"   [OK] {edits_applied_count} modificaciones y {new_added_count} pedidos nuevos integrados exitosamente.")
        except Exception as e_man:
            print(f"   [AVISO] No se pudo procesar {json_manual_path}: {e_man}")

    # 7. Generar Resumen de Cargas (Solo Pedidos VIGENTES - Excluye Finalizados)
    vigentes_activos = [s for s in active_shipments if s.get('estatus') != 'Finalizado']
    vigentes_flujo = [s for s in flow_orders if s.get('estatus') != 'Finalizado']
    finalizados_activos = [s for s in active_shipments if s.get('estatus') == 'Finalizado']
    finalizados_flujo = [s for s in flow_orders if s.get('estatus') == 'Finalizado']

    activos_count = Counter(s['responsable'] for s in vigentes_activos)
    flujo_count = Counter(s['responsable'] for s in vigentes_flujo)
    finalizados_count = Counter(s['responsable'] for s in (finalizados_activos + finalizados_flujo))
    all_analysts_set = sorted(list(set(list(activos_count.keys()) + list(flujo_count.keys()))))

    analysts_summary = []
    print("\n-------------------------------------------------------------------------------------------------------------------------------------------------")
    print(f"{'ANALISTA':<18} | {'ACTIVOS':<8} | {'POR FACTURAR':<12} | {'TOTAL VIG':<10} | {'FIN':<6} | {'TRÁNSITO':<18} | {'DISPONIBILIDAD APOYO'}")
    print("-------------------------------------------------------------------------------------------------------------------------------------------------")
    for a in all_analysts_set:
        if a in ['FRESCO_OMITIR', 'Marcelo']:
            continue
        act = activos_count[a]
        flj = flujo_count[a]
        tot = act + flj
        fin_c = finalizados_count.get(a, 0)

        # Reglas operativas de apoyo y tipo de tránsito
        if a == 'Eugenia':
            tipo_transito = "Tránsito Corto"
            puede_apoyar = True
            salvedad_apoyo = "Apoyo Condicionado (Tránsito Corto - Priorizar cartera)"
            rol_desc = "[!] Apoyo Condicionado (Tránsito Corto)"
        elif a in ['Terrestre', 'Sin Asignar']:
            tipo_transito = "No aplica"
            puede_apoyar = False
            salvedad_apoyo = "No aplica"
            rol_desc = "No aplica"
        else:
            tipo_transito = "Regular (+30/+60d)"
            puede_apoyar = True
            salvedad_apoyo = "Habilitado (Tránsito Holgado +30/+60d)"
            rol_desc = "[OK] Habilitado (Tránsito Holgado)"

        analysts_summary.append({
            "analyst": a,
            "activos": act,
            "por_zarpar": flj,
            "total_proyectado": tot,
            "finalizados": fin_c,
            "tipo_transito": tipo_transito,
            "apoyo_equipo": puede_apoyar,
            "salvedad_apoyo": salvedad_apoyo
        })
        print(f"{a:<18} | {act:<8} | {flj:<12} | {tot:<10} | {fin_c:<6} | {tipo_transito:<18} | {rol_desc}")
    print("-------------------------------------------------------------------------------------------------------------------------------------------------")

    # 7.5. Limpieza previa y normalización canónica de pedidos
    # Construir mapa canónico de nombres oficiales por código SAP
    canonical_name_by_code = {}
    for item in cartera_raw_list:
        c_code = str(item.get("codigo") or "").strip()
        if c_code and item.get("cliente"):
            canonical_name_by_code[c_code] = clean_client_name(item["cliente"])
    for r in CUSTOMER_RULES:
        c_code = str(r.get("codigo_sap") or "").strip()
        if c_code and (r.get("nombre_oficial") or r.get("cliente_patron")):
            canonical_name_by_code[c_code] = clean_client_name(r.get("nombre_oficial") or r.get("cliente_patron"))

    for s in active_shipments + flow_orders:
        # 1. Limpieza de nombre de cliente y descarte de sufijos de pedido previo
        c_orig = s.get('cliente', '')
        c_clean = clean_client_name(c_orig)
        s['cliente'] = c_clean

        # 2. Rescate de Código SAP canónico
        if not s.get('codigo_cliente') or str(s.get('codigo_cliente')).lower() in ('desconocido', 'nan', 'none', '-'):
            rescued_code = find_sap_code_for_client(c_clean)
            if rescued_code:
                s['codigo_cliente'] = rescued_code

        # 3. Homologación al nombre oficial canónico si tiene código SAP
        c_code = str(s.get('codigo_cliente') or '').strip()
        if c_code and c_code in canonical_name_by_code:
            s['cliente'] = canonical_name_by_code[c_code]

        # 4. Asegurar asignación a titular de cartera si el pedido estaba como Sin Asignar
        if c_code and c_code in codigo_portfolio and (not s.get('responsable') or s.get('responsable') == 'Sin Asignar'):
            s['responsable'] = codigo_portfolio[c_code]

        # 5. Condición de venta: Intercompany pasa a 'Crédito'
        if str(s.get('condicion_venta') or '').strip().lower() == 'intercompany':
            s['condicion_venta'] = 'Crédito'

        # 6. Saneamiento de Folios: El Folio SII (factura) nunca debe ser un folio SAP (comienza con 9 y len >= 8)
        fac_val = str(s.get('factura') or '').strip()
        doc_val = str(s.get('doc_factura') or '').strip()
        if fac_val.startswith('9') and len(fac_val) >= 8:
            if not doc_val:
                s['doc_factura'] = fac_val
            s['factura'] = ''

        # 7. Perfil y Reglas de Cliente (Intercompany vs Comercial/Crédito)
        rule = get_customer_rule(s.get('cliente'), s.get('codigo_cliente'))
        is_intercompany_cli = bool(
            (rule and rule.get('categoria') == 'Intercompany') or
            any(k in str(s.get('cliente', '')).upper() for k in ('ANDES ASIA', 'AQUACHILE INC', 'PRODUCTOS ALIMENTICIOS SUPER', 'AGROSUPER CHINA'))
        )
        cat = 'Intercompany' if is_intercompany_cli else (rule.get('categoria', 'Comercial') if rule else (s.get('categoria_cliente') or 'Comercial'))
        s['categoria_cliente'] = cat

        # 8. Formato de Liberación:
        fmt = re.sub(r'\s*\((?:intercompany|cliente interno)\)', '', str(s.get('formato_liberacion') or ''), flags=re.IGNORECASE).strip()
        if not fmt or fmt in ('-', 'None'):
            if rule and rule.get('formato_liberacion'):
                fmt = rule.get('formato_liberacion')
            else:
                fmt = 'SWB' if is_intercompany_cli else 'OBL'
        s['formato_liberacion'] = fmt

        # 9. Release / SWB: SOLO fechas reales de liberación (YYYY-MM-DD) o vacío ("")
        cur_swb = str(s.get('swb_liberacion') or s.get('swb') or s.get('release') or '').strip()
        swb_dt = parse_comex_date(cur_swb)
        valid_release_str = swb_dt.strftime('%Y-%m-%d') if swb_dt else ''
        s['swb_liberacion'] = valid_release_str
        s['swb'] = valid_release_str
        s['release'] = valid_release_str

    # 7.6. Generar Directorio Unificado de Cartera de Clientes y Coberturas
    pedidos_por_cliente = defaultdict(lambda: {"activos": 0, "facturados": 0, "finalizados": 0, "flujo": 0})
    pedidos_por_codigo = defaultdict(lambda: {"activos": 0, "facturados": 0, "finalizados": 0, "flujo": 0})

    for s in active_shipments:
        c_name = str(s.get('cliente', '')).strip().upper()
        c_code = str(s.get('codigo_cliente', '')).strip()
        is_fin = (s.get('estatus') == 'Finalizado')
        has_fac = bool(s.get('factura') or s.get('doc_factura'))

        if c_name:
            if is_fin: pedidos_por_cliente[c_name]["finalizados"] += 1
            else:
                pedidos_por_cliente[c_name]["activos"] += 1
                if has_fac: pedidos_por_cliente[c_name]["facturados"] += 1
        if c_code:
            if is_fin: pedidos_por_codigo[c_code]["finalizados"] += 1
            else:
                pedidos_por_codigo[c_code]["activos"] += 1
                if has_fac: pedidos_por_codigo[c_code]["facturados"] += 1

    for s in flow_orders:
        c_name = str(s.get('cliente', '')).strip().upper()
        c_code = str(s.get('codigo_cliente', '')).strip()
        if c_name:
            pedidos_por_cliente[c_name]["flujo"] += 1
        if c_code:
            pedidos_por_codigo[c_code]["flujo"] += 1

    cartera_by_key = {}
    seen_codes = set()
    seen_names = set()

    # 1. Cargar clientes oficiales de la cartera (Asignacion Clientes.xlsx)
    for item in cartera_raw_list:
        c_name = clean_client_name(item["cliente"])
        code_val = str(item.get("codigo", "") or find_sap_code_for_client(c_name) or "").strip()
        key = ("CODE_" + code_val) if code_val else ("NAME_" + c_name.upper())

        rule = get_customer_rule(c_name, code_val)
        cat = rule.get("categoria", "Comercial") if rule else "Comercial"
        fmt_lib = rule.get("formato_liberacion", "OBL") if rule else "OBL"
        req_dhl = rule.get("requiere_dhl", True) if rule else True
        crit = rule.get("criterio_cierre", "general") if rule else "general"

        if code_val and code_val in pedidos_por_codigo:
            counts = pedidos_por_codigo[code_val]
        else:
            counts = pedidos_por_cliente.get(c_name.upper(), {"activos": 0, "facturados": 0, "finalizados": 0, "flujo": 0})
            if counts["activos"] == 0 and counts["finalizados"] == 0 and counts["flujo"] == 0:
                for k_cli, v_counts in pedidos_por_cliente.items():
                    if c_name.upper() in k_cli or (rule and rule.get("cliente_patron", "").upper() in k_cli):
                        counts = v_counts
                        break

        if key in cartera_by_key:
            existing = cartera_by_key[key]
            if existing["analista"] in ("Sin Asignar", "") and item.get("analista") not in ("Sin Asignar", ""):
                existing["analista"] = item["analista"]
            continue

        cartera_entry = {
            "codigo": code_val,
            "cliente": c_name,
            "analista": item["analista"],
            "contacto": item.get("contacto", ""),
            "celular": item.get("celular", ""),
            "categoria": cat,
            "formato_liberacion": fmt_lib,
            "requiere_dhl": req_dhl,
            "criterio_cierre": crit,
            "pedidos_activos": counts["activos"],
            "pedidos_facturados": counts["facturados"],
            "pedidos_finalizados": counts["finalizados"],
            "pedidos_flujo": counts["flujo"],
            "total_pedidos": counts["activos"] + counts["finalizados"] + counts["flujo"]
        }
        cartera_by_key[key] = cartera_entry
        seen_names.add(c_name.upper())
        if code_val:
            seen_codes.add(code_val)

    # 2. Incorporar clientes que tengan pedidos en la base pero no estuvieran en la planilla de cartera oficial
    for s in active_shipments + flow_orders:
        c_orig = s.get('cliente', '')
        c_name = clean_client_name(c_orig)
        code_val = str(s.get('codigo_cliente') or find_sap_code_for_client(c_name) or '').strip()
        key = ("CODE_" + code_val) if code_val else ("NAME_" + c_name.upper())

        if key in cartera_by_key or (code_val and code_val in seen_codes) or (c_name.upper() in seen_names):
            continue

        if code_val and code_val in pedidos_por_codigo:
            counts = pedidos_por_codigo[code_val]
        else:
            counts = pedidos_por_cliente.get(c_name.upper(), {"activos": 0, "facturados": 0, "finalizados": 0, "flujo": 0})

        if counts["activos"] == 0 and counts["finalizados"] == 0 and counts["flujo"] == 0:
            continue

        rule = get_customer_rule(c_name, code_val)
        ana_orig = s.get('responsable') or 'Sin Asignar'
        if ana_orig == 'Sin Asignar' and code_val and code_val in codigo_portfolio:
            ana_orig = codigo_portfolio[code_val]

        cat = rule.get("categoria", "Comercial") if rule else "Comercial"
        fmt_lib = rule.get("formato_liberacion", s.get("formato_liberacion", "OBL")) if rule else (s.get("formato_liberacion") or "OBL")

        cartera_entry = {
            "codigo": code_val,
            "cliente": c_name,
            "analista": ana_orig,
            "contacto": "",
            "celular": "",
            "categoria": cat,
            "formato_liberacion": fmt_lib,
            "requiere_dhl": rule.get("requiere_dhl", True) if rule else True,
            "criterio_cierre": rule.get("criterio_cierre", "general") if rule else "general",
            "pedidos_activos": counts["activos"],
            "pedidos_facturados": counts["facturados"],
            "pedidos_finalizados": counts["finalizados"],
            "pedidos_flujo": counts["flujo"],
            "total_pedidos": counts["activos"] + counts["finalizados"] + counts["flujo"]
        }
        cartera_by_key[key] = cartera_entry
        seen_names.add(c_name.upper())
        if code_val:
            seen_codes.add(code_val)

    clientes_cartera_list = list(cartera_by_key.values())
    clientes_cartera_list.sort(key=lambda x: (x["analista"], x["cliente"]))

    # Filtrado estricto final: eliminar cualquier pedido del equipo Fresco
    active_shipments = [o for o in active_shipments if o.get('responsable') != 'FRESCO_OMITIR' and o.get('responsable') != 'Fresco']
    flow_orders = [o for o in flow_orders if o.get('responsable') != 'FRESCO_OMITIR' and o.get('responsable') != 'Fresco']

    js_data = {
        "activeShipments": active_shipments,
        "flowOrders": flow_orders,
        "analystsSummary": analysts_summary,
        "clientesCartera": clientes_cartera_list,
        "reglasClientes": CUSTOMER_RULES
    }

    js_content = "window.CONGELADO_APP_DATA = " + json.dumps(js_data, ensure_ascii=False, indent=2) + ";\n"
    js_content += "window.INITIAL_CONGELADO_DATA = window.CONGELADO_APP_DATA.activeShipments;\n"
    js_content += "window.INITIAL_FLUJO_DATA = window.CONGELADO_APP_DATA.flowOrders;\n"
    js_content += "window.ANALYSTS_SUMMARY = window.CONGELADO_APP_DATA.analystsSummary;\n"
    js_content += "window.CLIENTES_CARTERA = window.CONGELADO_APP_DATA.clientesCartera;\n"
    js_content += "window.REGLAS_CLIENTES = window.CONGELADO_APP_DATA.reglasClientes;\n"

    tmp_path = "data_control_congelado.js.tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        f.write(js_content)
    os.replace(tmp_path, "data_control_congelado.js")

    print("\n>> data_control_congelado.js actualizado con exito con separacion de equipos Fresco / Congelado!")

if __name__ == '__main__':
    main()
